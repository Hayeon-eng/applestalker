#!/usr/bin/env python3
"""
hc_import.py — honeyComb 🐝C · 기존 수작업 산출물 → 목업 데이터셋 [2026-09 신규]

입력
  --report  Google_Shopping_GMC_Final_Report.xlsx  (Summary: 국가×제품 S.com 노출 O/X·1st exposed store,
                                                    Shopping Check: 국가×제품×속성 Entered/Missed)
  --master  __FFF8_GMC_Monitoring_Master_File_*.xlsx (Standard: 51 속성 표준, 스토어 시트: 주차별 피드 입력값)
출력
  hc_mock_data.json — HoneyComb 엔진/화면이 쓰는 표준 스키마(실수집 provider 로 바꿔도 같은 스키마):
    config.countries / config.products / config.keywords / config.top_n
    attributes[51] {no, category, name, code, observe: card|detail|feed}
    runs[] {run_id, week, at, source, cells[] {country, product, keyword, position, scom_exposed,
            first_store, status, attrs{code: entered|missed|na}, feed{code: fed|empty|na}}}

목업 규칙(수집 없음 — Final Report 의 수동 결과를 위치로 환산)
  · 1st exposed brand store == Samsung 이고 S.com O → position 1
  · S.com O 인데 1위가 다른 판매처 → position 3(상단 이내, 1위 아님) — 실제 위치는 리포트에 없음
  · S.com X → position None(미노출) · 리포트 공란(워치) → status "unchecked"
  · 자연어 long-tail 키워드는 목업에서 "unchecked"(수집 전) 로 둔다
"""
from __future__ import annotations
import argparse
import json
import re
from datetime import datetime
from typing import Any, Dict, List

COUNTRIES = [  # Final Report 국가명 ↔ 큐비 사이트코드 ↔ Google gl/hl ↔ Master File 스토어 시트
    {"code": "US", "label": "United States", "sitecode": "us", "gl": "us", "hl": "en", "storefront": "SEA", "currency": "USD"},
    {"code": "UK", "label": "United Kingdom", "sitecode": "uk", "gl": "uk", "hl": "en", "storefront": "SEUK", "currency": "GBP"},
    {"code": "DE", "label": "Germany", "sitecode": "de", "gl": "de", "hl": "de", "storefront": "SEG", "currency": "EUR"},
    {"code": "SG", "label": "Singapore", "sitecode": "sg", "gl": "sg", "hl": "en", "storefront": "SESP", "currency": "SGD"},
    {"code": "IN", "label": "India", "sitecode": "in", "gl": "in", "hl": "en", "storefront": "SIEL", "currency": "INR"},
    {"code": "AU", "label": "Australia", "sitecode": "au", "gl": "au", "hl": "en", "storefront": "SEAU", "currency": "AUD"},
]
REPORT_COUNTRY = {"US": "US", "UK": "UK", "Germany": "DE", "Singapore": "SG", "India": "IN", "Australia": "AU"}

PRODUCTS = [  # Master File 열 코드 ↔ Final Report 제품명 ↔ 큐비 슬러그
    {"slug": "galaxy-z-fold8-ultra", "label": "Galaxy Z Fold8 Ultra", "master_col": "Q8", "report_names": ["Fold8 Ultra"]},
    {"slug": "galaxy-z-fold8", "label": "Galaxy Z Fold8", "master_col": "H8", "report_names": ["Fold8"]},
    {"slug": "galaxy-z-flip8", "label": "Galaxy Z Flip8", "master_col": "B8", "report_names": ["Flip8"]},
    {"slug": "galaxy-watch-ultra2", "label": "Galaxy Watch Ultra2", "master_col": "V2", "report_names": ["Watch Ultra 2", "WatchUltra2"]},
    {"slug": "galaxy-watch9", "label": "Galaxy Watch9", "master_col": "Fresh9 (40mm)", "report_names": ["Watch 9", "Watch9"]},
]

# 키워드 초안 — brand(제품명) + long-tail(자연어). 운영 시 hc_config 에서 편집.
KEYWORDS = {
    "galaxy-z-fold8-ultra": [("Galaxy Z Fold8 Ultra", "brand"), ("Samsung Fold8 Ultra 512GB price", "longtail"), ("best foldable phone for multitasking", "longtail")],
    "galaxy-z-fold8": [("Galaxy Z Fold8", "brand"), ("Samsung Galaxy Z Fold8 256GB", "longtail"), ("lightest foldable phone 2026", "longtail")],
    "galaxy-z-flip8": [("Galaxy Z Flip8", "brand"), ("Samsung Flip8 pink", "longtail"), ("compact flip phone with good camera", "longtail")],
    "galaxy-watch-ultra2": [("Galaxy Watch Ultra2", "brand"), ("Samsung Watch Ultra 2 titanium", "longtail"), ("rugged smartwatch for diving", "longtail")],
    "galaxy-watch9": [("Galaxy Watch9", "brand"), ("Samsung Galaxy Watch9 40mm", "longtail"), ("best health tracking smartwatch", "longtail")],
}

# 51 속성의 관측 가능 위치(어디서 역추적하는가). 화면 카드/상세패널에서 안 보이는 것은 feed(피드 전용 — Merchant API·GMC 화면).
OBSERVE = {
    "card": {"title", "price", "sale_price", "image_link", "availability", "shipping", "installment", "brand", "condition",
             "free_shipping_threshold"},
    "detail": {"description", "additional_image_link", "video_link", "virtual_model_link", "product_highlight", "product_detail",
               "size", "color", "material", "variant_option", "question_and_answer", "related_product", "document_link",
               "energy_efficiency_class", "product_weight", "product_length", "product_width", "product_height", "gender",
               "google_product_category", "product_type", "gtin", "mpn", "item_group_title", "short_title", "lifestyle_image_link",
               "popularity_rank", "certification", "link", "mobile_link", "availability_date", "sale_price_effective_date",
               "expiration_date", "promotion_id"},
}


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def load_standard(master_path: str) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(master_path, read_only=True, data_only=True)
    ws = wb["Standard"]
    out, last_no = [], 0
    for r in list(ws.iter_rows(values_only=True))[2:]:
        if not r or not r[2] or not r[3]:
            continue
        name, code = _norm(r[2]), _norm(r[3])
        if not re.match(r"^[a-z_0-9]+$", code):
            continue
        no = int(r[0]) if r[0] and str(r[0]).strip().isdigit() else None
        if no:
            last_no = no
        observe = "card" if code in OBSERVE["card"] else "detail" if code in OBSERVE["detail"] else "feed"
        out.append({"no": no or last_no, "sub": no is None, "category": _norm(r[1]), "name": name, "code": code,
                    "feeding": _norm(r[4]), "conversational": name.endswith("*"), "observe": observe})
    return out


def load_feed_status(master_path: str, attributes: List[Dict[str, Any]]) -> Dict[str, Dict[str, Dict[str, Dict[str, str]]]]:
    """→ feed[week][country_code][product_slug][attr_name] = fed|empty"""
    from openpyxl import load_workbook
    wb = load_workbook(master_path, read_only=True, data_only=True)
    feed: Dict[str, Dict[str, Dict[str, Dict[str, str]]]] = {"Week0": {}, "Week1": {}}
    col_by_product = {p["master_col"]: p["slug"] for p in PRODUCTS}
    for c in COUNTRIES:
        if c["storefront"] not in wb.sheetnames:
            continue
        rows = list(wb[c["storefront"]].iter_rows(values_only=True))
        header = rows[4]
        # 헤더에서 제품 열 위치: 처음 5개(Week0), 다음 5개(Week1)
        idxs = [i for i, h in enumerate(header) if h in col_by_product]
        wk = {"Week0": idxs[0:5], "Week1": idxs[5:10]}
        for r in rows[5:]:
            if not r or not r[1]:
                continue
            name = _norm(r[1])
            for week, cols in wk.items():
                for i in cols:
                    slug = col_by_product[header[i]]
                    v = r[i] if i < len(r) else None
                    feed[week].setdefault(c["code"], {}).setdefault(slug, {})[name] = \
                        "fed" if (v is not None and _norm(v) not in ("", "-", "TBU", "X", "N/A")) else "empty"
    return feed


def load_report(report_path: str):
    """Summary → (country, product) → {scom, first_store}; Shopping Check → (country, product, attr) → entered"""
    from openpyxl import load_workbook
    wb = load_workbook(report_path, read_only=True, data_only=True)
    expo: Dict[tuple, Dict[str, Any]] = {}
    current_product = None
    for r in wb["Summary"].iter_rows(values_only=True):
        cells = [c for c in r if c is not None]
        if not cells:
            continue
        if cells[0] == "Country" and len(cells) > 1:
            current_product = _norm(cells[1]); continue
        cc = REPORT_COUNTRY.get(_norm(cells[0]))
        if cc and current_product:
            scom = next((str(c).strip() for c in r[2:4] if str(c or "").strip() in ("O", "X")), None)
            store = None
            for c in r[2:5]:
                s = _norm(c)
                if s and s not in ("O", "X", "Link") and not s.startswith("Entered") and not s.startswith("Missed"):
                    store = s
            expo[(cc, current_product)] = {"scom": scom, "first_store": store}
    entered: Dict[tuple, str] = {}
    for r in list(wb["Shopping Check"].iter_rows(values_only=True))[1:]:
        cc = REPORT_COUNTRY.get(_norm(r[0]))
        if cc:
            entered[(cc, _norm(r[1]), _norm(r[2]))] = "entered" if _norm(r[3]) == "O" else "missed"
    return expo, entered


def build(report_path: str, master_path: str, top_n: int = 8) -> Dict[str, Any]:
    attributes = load_standard(master_path)
    feed = load_feed_status(master_path, attributes)
    expo, entered = load_report(report_path)
    names_by_slug = {p["slug"]: p["report_names"] for p in PRODUCTS}

    def attrs_for(cc: str, slug: str) -> Dict[str, str]:
        out = {}
        for a in attributes:
            v = None
            for rn in names_by_slug[slug]:
                v = entered.get((cc, rn, a["name"])) or entered.get((cc, rn, a["name"].rstrip("*")))
                if v:
                    break
            out[a["code"] + ("" if not a["sub"] else f"#{a['no']}")] = v or ("na" if a["observe"] == "feed" else "missed")
        return out

    def cell(cc: str, slug: str, kw: str, ktype: str, week: str) -> Dict[str, Any]:
        e = None
        for rn in names_by_slug[slug]:
            e = expo.get((cc, rn))
            if e:
                break
        pos, status = None, "unchecked"
        if ktype == "brand" and e and e.get("scom") in ("O", "X"):
            if e["scom"] == "O" and (e.get("first_store") or "").lower().startswith("samsung"):
                pos, status = 1, "top1"
            elif e["scom"] == "O":
                pos, status = 3, "topn"
            else:
                pos, status = None, "absent"
        f = feed.get(week, {}).get(cc, {}).get(slug, {})
        return {"country": cc, "product": slug, "keyword": kw, "keyword_type": ktype,
                "position": pos, "top_n": top_n, "scom_exposed": (e or {}).get("scom"),
                "first_store": (e or {}).get("first_store"), "status": status,
                "attrs": attrs_for(cc, slug) if status != "unchecked" else {},
                "feed": {a["code"] + ("" if not a["sub"] else f"#{a['no']}"): f.get(a["name"], "na") for a in attributes},
                "evidence": None}

    runs = []
    for week, at in (("Week0", "2026-07-23"), ("Week1", "2026-07-27")):
        cells = [cell(c["code"], p["slug"], kw, kt, week) for c in COUNTRIES for p in PRODUCTS for kw, kt in KEYWORDS[p["slug"]]]
        runs.append({"run_id": f"hc_{week.lower()}_{at.replace('-', '')}", "week": week, "at": at,
                     "source": "mock:Final_Report+Master_File", "cells": cells})
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "config": {"top_n": top_n, "countries": COUNTRIES, "products": PRODUCTS,
                   "keywords": {k: [{"text": t, "type": ty} for t, ty in v] for k, v in KEYWORDS.items()},
                   "status_legend": {"top1": "우리 제품이 1위(position 1)", "topn": f"상단({top_n}위 안)이지만 1위는 다른 판매처",
                                     "low": f"노출은 되지만 {top_n}위 밖", "absent": "검색 결과에 우리 제품이 없음", "unchecked": "아직 조회하지 않음"}},
        "attributes": attributes, "runs": runs,
    }


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--report", required=True)
    ap.add_argument("--master", required=True)
    ap.add_argument("--out", default="hc_mock_data.json")
    ap.add_argument("--top-n", type=int, default=8)
    a = ap.parse_args()
    d = build(a.report, a.master, a.top_n)
    json.dump(d, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    cells = d["runs"][-1]["cells"]
    import collections
    print("attributes", len(d["attributes"]), "| cells", len(cells), collections.Counter(c["status"] for c in cells))
