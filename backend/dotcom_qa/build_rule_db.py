"""
build_rule_db.py — 큐비 — Fold7/Flip7 Spec Rule DB 생성기 [V3]

검증된 삼성 공식 스펙(samsung.com specs 페이지 기준, 2026-07 확인)과
다국어 Canonical Dictionary를 담은 Rule DB 엑셀 + 시드 JSON을 생성한다.

실행:  python build_rule_db.py
산출:  QA_Bee_Rule_DB.galaxy-z-fold7.xlsx / QA_Bee_Rule_DB.galaxy-z-flip7.xlsx
       spec_rules.seed.galaxy-z-fold7.json / spec_rules.seed.galaxy-z-flip7.json

기준값 출처(공식):
  · Fold7: Typical 4400 / Rated 4272 mAh · 215g · 8.9/4.2mm · 8.0" 2184x1968 ·
    Cover 6.5" 2520x1080 · 2600nits · SD 8 Elite · 12/16GB · 256/512GB/1TB ·
    200+12+10MP · front/cover 10MP · 25W · Wi-Fi 7 · IP48 · Android 16 · 영상 24h
  · Flip7: Typical 4300 / Rated 4174 mAh · 188g · 13.7/6.5mm · 6.9" 2520x1080 ·
    Cover 4.1" 1048x948 · 2600nits · Exynos 2500 · 12GB · 256/512GB ·
    50+12MP · front 10MP · 25W · Wi-Fi 7 · IP48 · Android 16 · 영상 31h

Dictionary 철학: 사람이 추가하거나 Gemini가 볼 여지를 최소화 —
삼성닷컴 주요 로케일의 표준 스펙 표기를 최대한 선탑재한다.
"""
from __future__ import annotations
import json
import sys

# ═══════════════════════════════════════════════════════════════════
# Canonical Dictionary — 대표 항목 → 다국어/동의어 Alias
#   엔진 규칙: 대표어가 attribute의 부분어면 alias 상속
#   (예: 'Storage' 사전은 'Storage Option'에도 적용)
# ═══════════════════════════════════════════════════════════════════
DICTIONARY = {
    "Weight": [
        "Device Weight", "Weight (g)",
        "무게", "중량",                          # ko
        "重量", "本体重量",                       # ja / zh
        "重量 (g)", "机身重量",                   # zh-CN
        "Gewicht",                               # de / nl
        "Poids",                                 # fr
        "Peso",                                  # es / it / pt
        "Waga",                                  # pl
        "Hmotnost",                              # cs
        "Súly",                                  # hu
        "Greutate",                              # ro
        "Βάρος",                                 # el
        "Ağırlık",                               # tr
        "Вес", "Вага",                           # ru / uk
        "الوزن",                                 # ar
        "משקל",                                  # he
        "น้ำหนัก",                                # th
        "Trọng lượng",                           # vi
        "Berat",                                 # id / ms
        "वज़न", "वजन",                            # hi
        "Vikt", "Vægt", "Vekt", "Paino",         # sv / da / no / fi
    ],
    "Main Display Size": [
        "Main Screen Size", "Main Display", "Main Screen", "Screen Size", "Display Size",
        "메인 디스플레이 크기", "메인 화면 크기", "화면 크기",
        "メインディスプレイサイズ", "メイン画面サイズ", "画面サイズ",
        "主屏尺寸", "主屏幕尺寸", "螢幕尺寸", "主螢幕尺寸",
        "Hauptdisplay-Größe", "Größe des Hauptdisplays", "Displaygröße", "Bildschirmgröße",
        "Taille de l'écran principal", "Taille de l'écran",
        "Tamaño de pantalla principal", "Tamaño de pantalla",
        "Dimensioni schermo principale", "Dimensioni display",
        "Tamanho da tela principal", "Tamanho do ecrã",
        "Schermgrootte", "Rozmiar ekranu głównego", "Rozmiar wyświetlacza",
        "Ana Ekran Boyutu", "Ekran Boyutu",
        "Размер основного экрана", "Размер экрана",
        "حجم الشاشة الرئيسية", "حجم الشاشة",
        "ขนาดหน้าจอหลัก", "ขนาดหน้าจอ",
        "Kích thước màn hình chính", "Kích thước màn hình",
        "Ukuran layar utama", "Ukuran layar",
    ],
    "Cover Display Size": [
        "Cover Screen Size", "Cover Screen", "Cover Display", "FlexWindow",
        "커버 디스플레이 크기", "커버 화면",
        "カバーディスプレイサイズ", "カバー画面",
        "外屏尺寸", "封面螢幕尺寸",
        "Cover-Display-Größe", "Frontdisplay",
        "Taille de l'écran extérieur", "Écran extérieur",
        "Pantalla frontal", "Pantalla externa",
        "Schermo esterno", "Tela externa", "Ecrã exterior",
        "Dış Ekran", "Внешний экран",
        "الشاشة الخارجية", "หน้าจอด้านนอก", "Màn hình ngoài", "Layar luar",
    ],
    "Resolution": [
        "Main Resolution", "Display Resolution", "Screen Resolution",
        "해상도", "解像度", "分辨率", "解析度",
        "Auflösung", "Résolution", "Resolución", "Risoluzione", "Resolução",
        "Resolutie", "Rozdzielczość", "Rozlišení", "Felbontás", "Rezoluție",
        "Ανάλυση", "Çözünürlük", "Разрешение", "Роздільна здатність",
        "الدقة", "دقة الشاشة", "רזולוציה",
        "ความละเอียด", "Độ phân giải", "Resolusi",
        "Upplösning", "Opløsning", "Oppløsning", "Tarkkuus",
    ],
    "Refresh Rate": [
        "주사율", "리프레시 레이트",
        "リフレッシュレート", "刷新率", "更新率", "更新頻率",
        "Bildwiederholrate", "Bildwiederholfrequenz",
        "Taux de rafraîchissement",
        "Frecuencia de actualización", "Tasa de refresco",
        "Frequenza di aggiornamento", "Taxa de atualização",
        "Verversingssnelheid", "Częstotliwość odświeżania", "Obnovovací frekvence",
        "Yenileme Hızı", "Частота обновления", "Частота оновлення",
        "معدل التحديث", "อัตรารีเฟรช", "Tần số quét", "Kecepatan refresh",
    ],
    "Peak Brightness": [
        "Max Brightness", "Brightness", "Maximum Brightness",
        "최대 밝기", "밝기", "ピーク輝度", "輝度", "峰值亮度", "亮度", "最大亮度",
        "Maximale Helligkeit", "Helligkeit", "Luminosité maximale", "Luminosité",
        "Brillo máximo", "Brillo", "Luminosità di picco", "Luminosità",
        "Brilho máximo", "Brilho", "Maksymalna jasność", "Jasność",
        "Maksimum Parlaklık", "Parlaklık", "Пиковая яркость", "Яркость",
        "السطوع", "ความสว่างสูงสุด", "Độ sáng tối đa", "Kecerahan puncak",
    ],
    "Processor": [
        "Chipset", "CPU", "AP", "Mobile Processor", "SoC",
        "프로세서", "칩셋", "プロセッサ", "チップセット", "处理器", "處理器", "芯片",
        "Prozessor", "Processeur", "Procesador", "Processore", "Processador",
        "Processor", "Procesor", "İşlemci", "Процессор", "Процесор",
        "المعالج", "מעבד", "โปรเซสเซอร์", "หน่วยประมวลผล", "Bộ xử lý", "Prosesor",
    ],
    "RAM": [
        "Memory", "Memory (RAM)", "RAM Size",
        "메모리", "램", "メモリ", "RAM容量", "内存", "運行內存", "記憶體",
        "Arbeitsspeicher", "Mémoire", "Mémoire vive",
        "Memoria RAM", "Memória RAM", "Geheugen", "Pamięć RAM", "Operační paměť",
        "Bellek", "Оперативная память", "Оперативна пам'ять",
        "ذاكرة الوصول العشوائي", "الذاكرة العشوائية",
        "หน่วยความจำ", "Bộ nhớ RAM", "Memori",
    ],
    "Storage": [
        "Internal Storage", "Internal Memory", "Storage Option", "Storage Options", "ROM",
        "저장 용량", "저장공간", "내장 메모리",
        "ストレージ", "内蔵ストレージ", "存储", "存储容量", "儲存空間", "內置存儲",
        "Speicher", "Interner Speicher", "Speicherplatz",
        "Stockage", "Mémoire de stockage",
        "Almacenamiento", "Memoria interna", "Archiviazione",
        "Armazenamento", "Memória interna", "Opslag", "Opslagruimte",
        "Pamięć wewnętrzna", "Vnitřní paměť", "Depolama", "Dahili Depolama",
        "Встроенная память", "Память", "Вбудована пам'ять",
        "التخزين", "سعة التخزين", "الذاكرة الداخلية",
        "ที่เก็บข้อมูล", "หน่วยความจำภายใน", "Bộ nhớ trong", "Penyimpanan", "Memori internal",
    ],
    "Rear Camera": [
        "Main Camera", "Back Camera", "Rear Camera Resolution", "Camera",
        "후면 카메라", "메인 카메라", "背面カメラ", "リアカメラ", "メインカメラ",
        "后置摄像头", "主摄像头", "後置鏡頭", "主鏡頭",
        "Rückkamera", "Hauptkamera", "Caméra arrière", "Caméra principale",
        "Cámara trasera", "Cámara principal", "Fotocamera posteriore", "Fotocamera principale",
        "Câmera traseira", "Câmera principal", "Câmara traseira",
        "Achtercamera", "Hoofdcamera", "Aparat tylny", "Aparat główny", "Zadní fotoaparát",
        "Arka Kamera", "Ana Kamera", "Основная камера", "Задняя камера",
        "الكاميرا الخلفية", "الكاميرا الرئيسية",
        "กล้องหลัง", "Camera sau", "Kamera belakang", "Kamera utama",
    ],
    "Front Camera": [
        "Selfie Camera", "Front Camera Resolution",
        "전면 카메라", "셀피 카메라", "フロントカメラ", "前面カメラ", "インカメラ",
        "前置摄像头", "前置鏡頭", "自拍鏡頭",
        "Frontkamera", "Caméra avant", "Caméra frontale", "Caméra selfie",
        "Cámara frontal", "Cámara para selfies", "Fotocamera anteriore", "Fotocamera frontale",
        "Câmera frontal", "Câmara frontal", "Voorcamera", "Aparat przedni", "Přední fotoaparát",
        "Ön Kamera", "Фронтальная камера", "Селфи-камера",
        "الكاميرا الأمامية", "กล้องหน้า", "Camera trước", "Kamera depan",
    ],
    "Cover Camera": [
        "Cover Screen Camera", "External Camera",
        "커버 카메라", "커버 스크린 카메라", "カバーカメラ", "外屏摄像头",
        "Cover-Kamera", "Caméra écran extérieur", "Cámara de pantalla frontal",
        "Fotocamera schermo esterno", "Câmera da tela externa",
        "Dış Ekran Kamerası", "Камера внешнего экрана", "กล้องหน้าจอด้านนอก",
    ],
    "Typical Battery Capacity": [
        "Typical Capacity", "Battery Capacity (Typical)", "Battery (Typical)",
        "Battery Capacity", "Battery",
        "배터리 용량(일반)", "배터리 용량 (일반)", "배터리 용량", "일반 배터리 용량",
        "バッテリー容量(標準)", "バッテリー容量（標準）", "バッテリー容量", "標準容量",
        "电池容量(典型值)", "电池容量（典型值）", "电池容量", "典型容量", "典型值",
        "電池容量(典型值)", "電池容量",
        "Akkukapazität (typisch)", "Typische Akkukapazität", "Typische Batteriekapazität",
        "Akkukapazität", "Batteriekapazität",
        "Capacité de la batterie (typique)", "Capacité typique", "Capacité de la batterie",
        "Capacidad de batería (típica)", "Capacidad típica", "Capacidad de la batería",
        "Capacità della batteria (tipica)", "Capacità tipica", "Capacità della batteria",
        "Capacidade da bateria (típica)", "Capacidade típica", "Capacidade da bateria",
        "Batterijcapaciteit (typisch)", "Typische batterijcapaciteit", "Batterijcapaciteit",
        "Pojemność baterii (typowa)", "Typowa pojemność baterii", "Pojemność baterii",
        "Kapacita baterie (typická)", "Typická kapacita",
        "Pil Kapasitesi (Tipik)", "Tipik Pil Kapasitesi", "Pil Kapasitesi", "Batarya Kapasitesi",
        "Ёмкость аккумулятора (типовая)", "Типовая ёмкость", "Ёмкость аккумулятора",
        "سعة البطارية (النموذجية)", "السعة النموذجية", "سعة البطارية",
        "ความจุแบตเตอรี่ (ทั่วไป)", "ความจุแบตเตอรี่",
        "Dung lượng pin (điển hình)", "Dung lượng pin",
        "Kapasitas baterai (tipikal)", "Kapasitas baterai",
    ],
    "Rated Battery Capacity": [
        "Rated Capacity", "Battery Capacity (Rated)", "Rated Value",
        "정격 용량", "배터리 용량(정격)",
        "定格容量", "定格値",
        "额定容量", "額定容量", "额定值",
        "Nennkapazität", "Nominale Kapazität",
        "Capacité nominale", "Capacidad nominal", "Capacità nominale", "Capacidade nominal",
        "Nominale capaciteit", "Pojemność znamionowa", "Jmenovitá kapacita",
        "Nominal Kapasite", "Anma Kapasitesi",
        "Номинальная ёмкость", "Номінальна ємність",
        "السعة المقدرة", "السعة الاسمية",
        "ความจุที่กำหนด", "Dung lượng định mức", "Kapasitas terukur",
    ],
    "Wired Charging": [
        "Fast Charging", "Wired Charging Speed", "Charging",
        "유선 충전", "고속 충전", "有線充電", "急速充電", "有线充电", "快充",
        "Kabelgebundenes Laden", "Schnellladen",
        "Charge filaire", "Charge rapide",
        "Carga con cable", "Carga rápida", "Ricarica via cavo", "Ricarica rapida",
        "Carregamento com fio", "Carregamento rápido",
        "Bedraad opladen", "Snel opladen", "Ładowanie przewodowe", "Szybkie ładowanie",
        "Kablolu Şarj", "Hızlı Şarj", "Проводная зарядка", "Быстрая зарядка",
        "الشحن السلكي", "الشحن السريع", "ชาร์จแบบมีสาย", "Sạc có dây", "Pengisian daya kabel",
    ],
    "Video Playback": [
        "Video Playback Time", "Video Playback (Hours)",
        "동영상 재생", "동영상 재생 시간", "ビデオ再生", "ビデオ再生時間",
        "视频播放", "視頻播放", "影片播放時間",
        "Videowiedergabe", "Videowiedergabezeit",
        "Lecture vidéo", "Durée de lecture vidéo",
        "Reproducción de video", "Riproduzione video", "Reprodução de vídeo",
        "Video afspelen", "Odtwarzanie wideo", "Video Oynatma",
        "Воспроизведение видео", "تشغيل الفيديو",
        "การเล่นวิดีโอ", "Phát video", "Pemutaran video",
    ],
    "Wi-Fi": [
        "WiFi", "Wi Fi", "Wireless LAN", "WLAN", "Wi-Fi Version",
        "와이파이", "무선랜",
        "無線LAN", "无线局域网",
    ],
    "Bluetooth": [
        "Bluetooth Version", "BT",
        "블루투스", "ブルートゥース", "蓝牙", "藍牙",
    ],
    "Thickness (Folded)": [
        "Folded Thickness", "Thickness Folded", "Dimensions (Folded)", "Folded",
        "두께(접었을 때)", "접었을 때",
        "厚さ(折りたたみ時)", "折りたたみ時",
        "厚度(折叠)", "折叠时厚度", "摺疊時",
        "Dicke (gefaltet)", "Zusammengeklappt",
        "Épaisseur (plié)", "Plié",
        "Grosor (plegado)", "Plegado", "Spessore (piegato)", "Da piegato",
        "Espessura (dobrado)", "Dobrado",
        "Grubość (złożony)", "Kalınlık (Katlı)",
        "Толщина (в сложенном виде)", "السماكة (مطوي)",
        "ความหนา (พับ)", "Độ dày (gập)", "Ketebalan (terlipat)",
    ],
    "Thickness (Unfolded)": [
        "Unfolded Thickness", "Thickness Unfolded", "Dimensions (Unfolded)", "Unfolded",
        "두께(펼쳤을 때)", "펼쳤을 때",
        "厚さ(展開時)", "展開時",
        "厚度(展开)", "展开时厚度", "展開時",
        "Dicke (aufgeklappt)", "Aufgeklappt",
        "Épaisseur (déplié)", "Déplié",
        "Grosor (desplegado)", "Desplegado", "Spessore (aperto)", "Da aperto",
        "Espessura (desdobrado)", "Desdobrado",
        "Grubość (rozłożony)", "Kalınlık (Açık)",
        "Толщина (в разложенном виде)", "السماكة (مفتوح)",
        "ความหนา (กาง)", "Độ dày (mở)", "Ketebalan (terbuka)",
    ],
    "Operating System": [
        "OS", "OS Version",
        "운영체제", "OS(운영체제)",
        "オペレーティングシステム", "操作系统", "作業系統",
        "Betriebssystem", "Système d'exploitation",
        "Sistema operativo", "Sistema operacional", "Besturingssysteem",
        "System operacyjny", "Operační systém", "İşletim Sistemi",
        "Операционная система", "Операційна система",
        "نظام التشغيل", "ระบบปฏิบัติการ", "Hệ điều hành", "Sistem operasi",
    ],
    "Water Resistance": [
        "IP Rating", "Water and Dust Resistance", "Dust and Water Resistance", "IP48",
        "방수", "방수·방진", "防水", "防水防塵", "防水防尘",
        "Wasserbeständigkeit", "Wasser- und Staubschutz",
        "Résistance à l'eau", "Résistance à l'eau et à la poussière",
        "Resistencia al agua", "Resistenza all'acqua", "Resistência à água",
        "Waterbestendigheid", "Wodoodporność", "Suya Dayanıklılık",
        "Водонепроницаемость", "Защита от воды и пыли",
        "مقاومة الماء", "กันน้ำ", "Kháng nước", "Tahan air",
    ],
    "Optical Zoom": [
        "Optical Quality Zoom", "Telephoto Zoom", "Zoom",
        "광학 줌", "光学ズーム", "光学变焦", "光學變焦",
        "Optischer Zoom", "Zoom optique", "Zoom óptico", "Zoom ottico",
        "Optik Yakınlaştırma", "Оптический зум", "زوم بصري",
        "ซูมออปติคอล", "Zoom quang học", "Zoom optik",
    ],
}

# ═══════════════════════════════════════════════════════════════════
# MasterSpec — 제품별 룰 (검증된 공식 기준값)
# 필드: rule_id, category, attribute, expected, unit, validation,
#       priority, page, interaction, exception, fix_guide, notes
# ═══════════════════════════════════════════════════════════════════
_COMMON_EXC_BAT_TYP = "Do not compare with Rated"
_COMMON_EXC_BAT_RAT = "Do not compare with Typical"

FOLD7_RULES = [
    ("DEVICE_001", "Device", "Product Name", "Galaxy Z Fold7", "", "exact", "Critical", "PDP", "", "", "Use official product name 'Galaxy Z Fold7'", ""),
    ("DEVICE_002", "Device", "Model Code", "SM-F966", "", "prefix", "Critical", "PDP", "", "Regional SKU suffix allowed (B/N/U/W/0 + /DS)", "Verify regional suffix only — base must be SM-F966", ""),
    ("DISPLAY_001", "Display", "Main Display Size", "8.0", "inch", "numeric_exact", "High", "PDP/Compare", "", "", "Update to 8.0 inch (global master)", "Full rectangle 8.0\""),
    ("DISPLAY_002", "Display", "Main Resolution", "2184 x 1968", "px", "exact", "High", "PDP/Compare", "", "", "Update to 2184 x 1968 (QXGA+)", ""),
    ("DISPLAY_003", "Display", "Cover Display Size", "6.5", "inch", "numeric_exact", "High", "PDP/Compare", "", "", "Update to 6.5 inch", ""),
    ("DISPLAY_004", "Display", "Cover Resolution", "2520 x 1080", "px", "exact", "High", "PDP/Compare", "", "", "Update to 2520 x 1080 (HD+ 21:9)", ""),
    ("DISPLAY_005", "Display", "Refresh Rate", "120", "Hz", "numeric_exact", "Medium", "PDP", "", "Adaptive 1-120Hz — highest value must be 120", "Update to 120Hz", ""),
    ("DISPLAY_006", "Display", "Peak Brightness", "2600", "nits", "numeric_exact", "Medium", "PDP", "", "", "Update to 2600 nits", ""),
    ("CPU_001", "Processor", "Processor", "Snapdragon 8 Elite", "", "dictionary", "High", "PDP/Compare", "", "Full name 'Snapdragon 8 Elite for Galaxy' — localized suffix allowed", "Must contain 'Snapdragon 8 Elite'", ""),
    ("MEMORY_001", "Memory", "RAM", "12 / 16", "GB", "option_match", "Critical", "PDP", "Variant", "16GB paired with 1TB model only", "12GB base · 16GB on 1TB variant", ""),
    ("MEMORY_002", "Storage", "Storage Option", "256 / 512 / 1TB", "GB", "option_match", "Critical", "Buy Box", "Variant", "1TB availability may vary by market — check Country Exception", "All storage options must be selectable", ""),
    ("CAM_001", "Camera", "Rear Camera", "200MP + 12MP + 10MP", "", "exact", "High", "PDP/Compare", "", "", "200MP Wide + 12MP Ultra Wide + 10MP Telephoto", ""),
    ("CAM_002", "Camera", "Front Camera", "10", "MP", "numeric_exact", "Medium", "PDP", "", "", "Main Screen camera 10MP (100° wide angle)", ""),
    ("CAM_003", "Camera", "Cover Camera", "10", "MP", "numeric_exact", "Medium", "PDP", "", "", "Cover Screen camera 10MP", ""),
    ("CAM_004", "Camera", "Optical Zoom", "3", "x", "numeric_exact", "Low", "PDP", "", "", "3x optical zoom (telephoto)", ""),
    ("BATTERY_001", "Battery", "Typical Battery Capacity", "4400", "mAh", "numeric_exact", "Critical", "PDP", "", _COMMON_EXC_BAT_TYP, "Change to 4400 mAh (typical)", ""),
    ("BATTERY_002", "Battery", "Rated Battery Capacity", "4272", "mAh", "numeric_exact", "Critical", "Disclaimer", "Expand", _COMMON_EXC_BAT_RAT, "Change to 4272 mAh (rated) in disclaimer", "Samsung official footnote value"),
    ("BATTERY_003", "Battery", "Wired Charging", "25", "W", "numeric_exact", "Medium", "PDP", "", "", "25W wired charging", ""),
    ("BATTERY_004", "Battery", "Video Playback", "24", "hours", "numeric_exact", "High", "PDP", "", "", "Up to 24 hours video playback", "Longest on any Z Fold"),
    ("CONNECT_001", "Connectivity", "Wi-Fi", "Wi-Fi 7", "", "dictionary", "Medium", "PDP", "", "802.11be availability varies by market", "Normalize spelling to Wi-Fi 7", ""),
    ("CONNECT_002", "Connectivity", "Bluetooth", "5.4", "", "exact", "Low", "PDP", "", "", "Bluetooth 5.4", ""),
    ("DIM_001", "Dimension", "Weight", "215", "g", "numeric_exact", "High", "PDP/Compare", "Dropdown", "Weight may vary slightly by country (official: 215g)", "Change to 215g", ""),
    ("DIM_002", "Dimension", "Thickness (Folded)", "8.9", "mm", "numeric_exact", "High", "PDP/Compare", "Dropdown", "", "Folded 8.9mm", ""),
    ("DIM_003", "Dimension", "Thickness (Unfolded)", "4.2", "mm", "numeric_exact", "High", "PDP/Compare", "Dropdown", "", "Unfolded 4.2mm", ""),
    ("DUR_001", "Durability", "Water Resistance", "IP48", "", "exact", "Medium", "PDP", "", "", "IP48 water and dust resistance", ""),
    ("OS_001", "Software", "Operating System", "Android 16", "", "exists", "Low", "PDP", "", "OS version increases with updates — check launch OS only at launch window", "Android 16 (One UI 8) at launch", ""),
    ("AI_001", "AI", "Galaxy AI", "Supported", "", "exists", "Medium", "PDP", "", "", "Galaxy AI mention must be present", ""),
    ("PKG_001", "Package", "USB-C Cable", "Included", "", "exists", "Low", "What's in the box", "Accordion", "", "USB-C cable listed in box contents", ""),
    ("PKG_002", "Package", "Ejection Pin", "Included", "", "exists", "Low", "What's in the box", "Accordion", "", "Ejection pin listed in box contents", ""),
    ("DISC_001", "Disclaimer", "Adapter not included", "Present", "", "exists", "Medium", "Disclaimer", "Expand", "Charger sold separately — wording varies by market", "Add adapter disclaimer if missing", ""),
    ("DISC_002", "Disclaimer", "Typical value tested", "Present", "", "exists", "Medium", "Disclaimer", "Expand", "", "IEC 61960 typical-value measurement disclaimer must be present", ""),
]

FLIP7_RULES = [
    ("DEVICE_001", "Device", "Product Name", "Galaxy Z Flip7", "", "exact", "Critical", "PDP", "", "", "Use official product name 'Galaxy Z Flip7'", ""),
    ("DEVICE_002", "Device", "Model Code", "SM-F766", "", "prefix", "Critical", "PDP", "", "Regional SKU suffix allowed (B/N/U/W/0 + /DS)", "Verify regional suffix only — base must be SM-F766", ""),
    ("DISPLAY_001", "Display", "Main Display Size", "6.9", "inch", "numeric_exact", "High", "PDP/Compare", "", "", "Update to 6.9 inch (full rectangle)", "6.8\" accounting rounded corners"),
    ("DISPLAY_002", "Display", "Main Resolution", "2520 x 1080", "px", "exact", "High", "PDP/Compare", "", "", "Update to 2520 x 1080 (FHD+ 21:9)", ""),
    ("DISPLAY_003", "Display", "Cover Display Size", "4.1", "inch", "numeric_exact", "High", "PDP/Compare", "", "", "Update to 4.1 inch (FlexWindow)", ""),
    ("DISPLAY_004", "Display", "Cover Resolution", "1048 x 948", "px", "exact", "High", "PDP/Compare", "", "", "Update to 1048 x 948", ""),
    ("DISPLAY_005", "Display", "Refresh Rate", "120", "Hz", "numeric_exact", "Medium", "PDP", "", "Adaptive 1-120Hz — highest value must be 120", "Update to 120Hz (main & cover)", ""),
    ("DISPLAY_006", "Display", "Peak Brightness", "2600", "nits", "numeric_exact", "Medium", "PDP", "", "", "Update to 2600 nits", ""),
    ("CPU_001", "Processor", "Processor", "Exynos 2500", "", "dictionary", "High", "PDP/Compare", "", "", "Must contain 'Exynos 2500' (3nm)", "First Exynos on Z Flip"),
    ("MEMORY_001", "Memory", "RAM", "12", "GB", "numeric_exact", "Critical", "PDP", "", "", "12GB only", ""),
    ("MEMORY_002", "Storage", "Storage Option", "256 / 512", "GB", "option_match", "Critical", "Buy Box", "Variant", "No 1TB option on Flip7", "Both storage options must be selectable", ""),
    ("CAM_001", "Camera", "Rear Camera", "50MP + 12MP", "", "exact", "High", "PDP/Compare", "", "", "50MP Wide + 12MP Ultra Wide", ""),
    ("CAM_002", "Camera", "Front Camera", "10", "MP", "numeric_exact", "Medium", "PDP", "", "", "Main Screen camera 10MP", ""),
    ("BATTERY_001", "Battery", "Typical Battery Capacity", "4300", "mAh", "numeric_exact", "Critical", "PDP", "", _COMMON_EXC_BAT_TYP, "Change to 4300 mAh (typical)", ""),
    ("BATTERY_002", "Battery", "Rated Battery Capacity", "4174", "mAh", "numeric_exact", "Critical", "Disclaimer", "Expand", _COMMON_EXC_BAT_RAT, "Change to 4174 mAh (rated) in disclaimer", "Samsung official footnote value"),
    ("BATTERY_003", "Battery", "Wired Charging", "25", "W", "numeric_exact", "Medium", "PDP", "", "", "25W wired charging", ""),
    ("BATTERY_004", "Battery", "Video Playback", "31", "hours", "numeric_exact", "High", "PDP", "", "", "Up to 31 hours video playback", ""),
    ("CONNECT_001", "Connectivity", "Wi-Fi", "Wi-Fi 7", "", "dictionary", "Medium", "PDP", "", "802.11be availability varies by market", "Normalize spelling to Wi-Fi 7", "Samsung official footnote lists Wi-Fi 7"),
    ("CONNECT_002", "Connectivity", "Bluetooth", "5.4", "", "exact", "Low", "PDP", "", "", "Bluetooth 5.4", "Verify on official spec page per market"),
    ("DIM_001", "Dimension", "Weight", "188", "g", "numeric_exact", "High", "PDP/Compare", "Dropdown", "Weight may vary slightly by country (official: 188g)", "Change to 188g", ""),
    ("DIM_002", "Dimension", "Thickness (Folded)", "13.7", "mm", "numeric_exact", "High", "PDP/Compare", "Dropdown", "", "Folded 13.7mm", ""),
    ("DIM_003", "Dimension", "Thickness (Unfolded)", "6.5", "mm", "numeric_exact", "High", "PDP/Compare", "Dropdown", "", "Unfolded 6.5mm", ""),
    ("DUR_001", "Durability", "Water Resistance", "IP48", "", "exact", "Medium", "PDP", "", "", "IP48 water and dust resistance", ""),
    ("OS_001", "Software", "Operating System", "Android 16", "", "exists", "Low", "PDP", "", "OS version increases with updates", "Android 16 (One UI 8) at launch", ""),
    ("AI_001", "AI", "Galaxy AI", "Supported", "", "exists", "Medium", "PDP", "", "", "Galaxy AI mention must be present", ""),
    ("PKG_001", "Package", "USB-C Cable", "Included", "", "exists", "Low", "What's in the box", "Accordion", "", "USB-C cable listed in box contents", ""),
    ("PKG_002", "Package", "Ejection Pin", "Included", "", "exists", "Low", "What's in the box", "Accordion", "", "Ejection pin listed in box contents", ""),
    ("DISC_001", "Disclaimer", "Adapter not included", "Present", "", "exists", "Medium", "Disclaimer", "Expand", "Charger sold separately — wording varies by market", "Add adapter disclaimer if missing", ""),
    ("DISC_002", "Disclaimer", "Typical value tested", "Present", "", "exists", "Medium", "Disclaimer", "Expand", "", "IEC 61960 typical-value measurement disclaimer must be present", ""),
]

EXCEPTIONS = [
    ("BATTERY", "Typical and Rated battery capacities are different attributes — never compared to each other"),
    ("DISCLAIMER", "Promotion/offer/banner sections are excluded from spec validation"),
    ("RESOLUTION", "Main and Cover display specs are different attributes — never compared to each other"),
    ("ADAPTIVE_REFRESH", "Adaptive refresh (1-120Hz) — validate the maximum value only"),
]

INTERACTIONS = [
    ("Buy Box", "Iterate storage/color variants"),
    ("Compare", "Iterate model dropdown"),
    ("Accordion", "Expand all (What's in the box, spec groups)"),
    ("Spec", "Click 'See all specs' / expand full spec table"),
    ("Disclaimer", "Expand footnotes / legal text"),
]

COUNTRY_EXCEPTIONS = [
    # (country/sitecode, rule, action) — 예시 정책: 반영 전 정책 확정 필요(Notes 참조)
    ("sg", "DISC_001", "allow-different-wording"),
    ("eu", "Energy Label", "skip"),
]
_CE_NOTE = "Example policy rows — confirm actual per-country policy before enforcing"

_VAL_LABEL = {"exact": "Exact", "numeric_exact": "Numeric Exact", "prefix": "Prefix Match",
              "dictionary": "Dictionary", "option_match": "Option Match", "exists": "Exists"}


def build_xlsx(product: str, rules: list, path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    hfont = Font(bold=True, color="FFFFFF"); hfill = PatternFill("solid", fgColor="1B2A4A")

    def _hdr(ws):
        for c in ws[1]:
            c.font = hfont; c.fill = hfill; c.alignment = Alignment(horizontal="center")

    ws = wb.active; ws.title = f"{product}_MasterSpec_V3"
    ws.append(["Rule ID", "Category", "Attribute", "Official Value", "Unit", "Validation",
               "Priority", "Page", "Interaction", "Exception", "Fix Guide", "Notes"])
    _hdr(ws)
    for r in rules:
        row = list(r); row[5] = _VAL_LABEL.get(row[5], row[5])
        ws.append(row)

    ws2 = wb.create_sheet("Dictionary")
    ws2.append(["Representative", "Alias"]); _hdr(ws2)
    for rep, aliases in DICTIONARY.items():
        for a in aliases:
            ws2.append([rep, a])

    ws3 = wb.create_sheet("ExceptionRule")
    ws3.append(["Rule", "Description"]); _hdr(ws3)
    for r in EXCEPTIONS:
        ws3.append(list(r))

    ws4 = wb.create_sheet("InteractionRule")
    ws4.append(["Section", "Action"]); _hdr(ws4)
    for r in INTERACTIONS:
        ws4.append(list(r))

    ws5 = wb.create_sheet("CountryException")
    ws5.append(["Country", "Rule", "Action"]); _hdr(ws5)
    for r in COUNTRY_EXCEPTIONS:
        ws5.append(list(r))
    ws5.append(["", "", ""]); ws5.append(["#", _CE_NOTE, ""])

    for w in wb.worksheets:
        for col, width in zip("ABCDEFGHIJKL", [12, 12, 26, 24, 6, 13, 9, 14, 12, 40, 44, 34]):
            w.column_dimensions[col].width = width
    wb.save(path)
    return path


def main():
    sys.path.insert(0, ".")
    import spec_rule_db
    for product, rules in (("galaxy-z-fold7", FOLD7_RULES), ("galaxy-z-flip7", FLIP7_RULES)):
        xlsx_path = f"QA_Bee_Rule_DB.{product}.xlsx"
        build_xlsx(product, rules, xlsx_path)
        content = open(xlsx_path, "rb").read()
        rs = spec_rule_db.parse_xlsx(content, product=product, version="V3")
        seed = f"spec_rules.seed.{product}.json"
        json.dump(rs, open(seed, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        n_alias = sum(len(v) for v in rs["dictionary"].values())
        print(f"{product}: rules={len(rs['rules'])} dict_reps={len(rs['dictionary'])} "
              f"aliases={n_alias} → {xlsx_path}, {seed}")


if __name__ == "__main__":
    main()
