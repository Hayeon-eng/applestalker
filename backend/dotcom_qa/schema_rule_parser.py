"""
schema_rule_parser.py — 큐비 — Dotcom QA 체커 [Phase A]

큐비 🐝 — 풀네임 QA Bee, 줄여서 큐비. 닷컴 페이지를 붕붕 돌며 규칙대로 검수.

Samsung 'JSON-LD Schema Copydeck' 엑셀(schema_deck.xlsx 의 M3/M12 시트)을
기계가 검수에 쓸 수 있는 '규칙 JSON'으로 변환한다.

엑셀 시트 레이아웃(파악된 규약):
  - 블록 시작:  c1='Type', c2='JSON-LD Keyword'(헤더행) → 바로 다음 행이 블록 본문 시작
  - 본문 시작:  c1=<블록명>, c2='@context'
  - c2='@type' → c5 에 타입(콤마 구분 가능: "WebPage,ItemPage")
  - c2='@id'   → c5 에 @id 패턴({SITECODE} 치환자 포함)
  - c2='Property', c5='Value' → 이후 행은 속성 행
  - 속성 행: c2=<속성명>, (중첩은 c3/c4), 값은 c5, 조건부 주석은 c6(Remarks)
  - Product.hasPart 는 c2='hasPart' c3='@id' 로 시작, 이후 c3='@id' 행들이 추가 @id

주의: 사람이 만든 엑셀이라 서식 변주가 있을 수 있어 방어적으로 파싱한다.
산출 규칙 JSON 은 '사람이 한 번 검토'하는 중간 산출물이다(하이브리드 방식).
"""
from __future__ import annotations
import json
import re
from typing import Any, Dict, List, Optional

# 파싱에 쓰는 열 인덱스(0-base). 시트가 A열 비움 + B열부터 내용.
C_NAME = 1      # 'Type' / 블록명 / (FAQ 라벨)
C_KW = 2        # 'JSON-LD Keyword' / 속성명 / @context,@type,@id,Property
C_N3 = 3        # 중첩 키워드(@id, @type, 중첩 속성)
C_N4 = 4        # 더 깊은 중첩
C_VAL = 5       # 'Class' / 'Value' / 실제 값
C_REM = 6       # 'Remarks'(조건부 주석)


def _s(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _clean_block_name(raw: str) -> str:
    # 'VideoObject\n1' → 'VideoObject #1'
    raw = raw.replace("\r", " ").replace("\n", " ").strip()
    m = re.match(r"^(VideoObject)\s+(\d+)$", raw)
    if m:
        return f"{m.group(1)} #{m.group(2)}"
    return raw


def _slug_from_id(id_pattern: str) -> Optional[str]:
    """@id 의 #fragment 나 마지막 경로 조각을 식별 슬러그로."""
    if not id_pattern:
        return None
    if "#" in id_pattern:
        return "#" + id_pattern.rsplit("#", 1)[-1]
    tail = id_pattern.rstrip("/").rsplit("/", 1)[-1]
    return tail or None


_ENUM_HINTS = {"webpage", "itempage", "product", "buyaction", "criticreview", "listitem",
               "imageobject", "videoobject", "3dmodel", "faqpage", "breadcrumblist", "offer",
               "aggregaterating", "brand", "organization", "question", "answer"}


def _kind(v: str) -> str:
    """기대값 종류: url / enum / text."""
    s = (v or "").strip()
    low = s.lower()
    if "{sitecode}" in low or "[sitecode]" in low or "{lang-code}" in low or low.startswith("http") or low.startswith("//"):
        return "url"
    if low in ("schema.org", "https://schema.org"):
        return "enum"
    # 공백 없는 짧은 토큰(콤마 구분 타입 포함) 이면서 enum 후보 → enum
    toks = [t.strip().lower() for t in s.replace("\n", ",").split(",") if t.strip()]
    if toks and all((" " not in t and (t in _ENUM_HINTS or t[:1].isupper() or "{" not in t)) for t in toks) and len(s) < 40:
        return "enum"
    return "text"


def parse_sheet(rows: List[tuple]) -> List[Dict[str, Any]]:
    """한 시트(M3 또는 M12)의 행들을 블록 규칙 리스트로 변환."""
    blocks: List[Dict[str, Any]] = []
    cur: Optional[Dict[str, Any]] = None
    mode = None          # None | 'awaiting_context' | 'head' | 'props'
    last_prop = None

    def cell(r, idx):
        return _s(r[idx]) if len(r) > idx else ""

    for r in rows:
        c_name = cell(r, C_NAME)
        c_kw = cell(r, C_KW)
        c_n3 = cell(r, C_N3)
        c_val = cell(r, C_VAL)
        c_rem = cell(r, C_REM)

        # 블록 헤더 감지
        if c_name == "Type" and "Keyword" in c_kw:
            mode = "awaiting_context"
            continue

        # 블록 본문 시작(@context 행)
        if c_kw == "@context":
            cur = {
                "name": _clean_block_name(c_name) or "(무명)",
                "types": [],
                "id_pattern": None,
                "id_slug": None,
                "required_properties": [],
                "optional_properties": [],   # Remarks 로 '조건부/후속' 표시된 속성
                "property_notes": {},
                "haspart_ids": [],
                "expected_values": {},   # 속성 → {"value","kind","nested"}
                "conditional": (c_rem if c_rem and c_rem not in ("-", "–", "—") else None),   # "특정 국가 제외" 등
            }
            blocks.append(cur)
            mode = "head"
            last_prop = None
            if c_val:
                cur["expected_values"]["@context"] = {"value": c_val, "kind": _kind(c_val), "nested": None}
            continue

        if cur is None:
            continue

        if c_kw == "@type":
            cur["types"] = [t.strip() for t in c_val.replace("\n", ",").split(",") if t.strip()]
            continue

        if c_kw == "@id" and not cur["id_pattern"]:
            cur["id_pattern"] = c_val or None
            cur["id_slug"] = _slug_from_id(c_val)
            continue

        if c_kw == "Property" and c_val.lower() == "value":
            mode = "props"
            last_prop = None
            continue

        if mode == "props":
            # 최상위 속성(c2 에 속성명)
            if c_kw and c_kw not in ("@id", "@type", "Property", "@context"):
                prop = c_kw
                if prop not in cur["required_properties"]:
                    cur["required_properties"].append(prop)
                last_prop = prop
                # 기대값 캡처(중첩 @id/@type 이면 nested 기록)
                if c_val:
                    nested = c_n3 if c_n3 in ("@id", "@type") else None
                    cur["expected_values"][prop] = {"value": c_val, "kind": _kind(c_val), "nested": nested}
                # Remarks 가 달린 속성(예: review = 출시 후 핫픽스)은 선택 처리
                if c_rem and c_rem not in ("-", "–", "—"):
                    if prop not in cur["optional_properties"]:
                        cur["optional_properties"].append(prop)
                    cur["property_notes"][prop] = c_rem
                if prop == "hasPart" and c_n3 == "@id" and c_val:
                    cur["haspart_ids"].append(c_val)
            # hasPart 의 이어지는 @id 행(c2 비고 c3=@id)
            elif not c_kw and c_n3 == "@id" and c_val and last_prop == "hasPart":
                cur["haspart_ids"].append(c_val)

    # 후처리: FAQPage 는 속성이 'FAQ N' 라벨로 되어 있어 c2 기반 수집이 비므로 보정
    for b in blocks:
        if "FAQPage" in b["types"] and not b["required_properties"]:
            b["required_properties"] = ["mainEntity"]
    return blocks


def build_rules(xlsx_path: str, sheets=("M3", "M12")) -> Dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    out: Dict[str, Any] = {
        "source": xlsx_path.rsplit("/", 1)[-1],
        "sitecode_token": "{SITECODE}",
        "products": {},
    }
    for sheet in sheets:
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        blocks = parse_sheet(rows)
        out["products"][sheet] = {
            "expected_types": sorted({t for b in blocks for t in b["types"]}),
            "block_count": len(blocks),
            "blocks": blocks,
        }
    wb.close()
    return out


if __name__ == "__main__":
    import os
    import sys
    src = sys.argv[1] if len(sys.argv) > 1 else "/mnt/user-data/uploads/schema_deck.xlsx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "schema_rules.json"
    rules = build_rules(src)
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(rules, f, ensure_ascii=False, indent=2)
    # 제품별 분리 파일(미니파이) — 로더가 우선 사용, 업로드 부담 완화
    base = os.path.dirname(dst) or "."
    meta = {"source": rules.get("source"), "sitecode_token": rules.get("sitecode_token")}
    for prod, pd in rules["products"].items():
        obj = {**meta, **pd}
        with open(os.path.join(base, f"schema_rules.{prod}.json"), "w", encoding="utf-8") as f:
            json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
        print(f"  → schema_rules.{prod}.json (블록 {len(pd.get('blocks', []))})")
    for prod, d in rules["products"].items():
        print(f"[{prod}] blocks={d['block_count']} types={d['expected_types']}")
        for b in d["blocks"]:
            hp = f" hasPart={len(b['haspart_ids'])}" if b["haspart_ids"] else ""
            cond = " (조건부)" if b["conditional"] else ""
            print(f"   - {b['name']:16s} type={b['types']} id={b['id_slug']} props={len(b['required_properties'])}{hp}{cond}")
    print(f"\n→ {dst} 저장")
