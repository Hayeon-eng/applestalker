"""
Dashboard Export Service
Exports dashboard data as PPTX or PNG files.
"""

import os
import io
from datetime import datetime
from typing import List, Dict, Any, Optional
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RgbColor
from pptx.enum.text import PP_ALIGN


class DashboardExportService:
    """
    Export dashboard data to PPTX or PNG format.
    """

    def __init__(self, output_dir: str = "./exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    async def export_to_pptx(
        self,
        data: Dict[str, Any],
        include_sections: List[str] = None,
    ) -> str:
        """
        Export dashboard data to PowerPoint presentation.
        
        Args:
            data: Dashboard export data
            include_sections: ['summary', 'changes', 'insights', 'actions']
            
        Returns:
            Path to generated PPTX file
        """
        if include_sections is None:
            include_sections = ["summary", "changes", "insights", "actions"]

        prs = Presentation()
        
        # Title Slide
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        
        title.text = f"🍎 Apple Tracker Report"
        subtitle.text = f"{data['site_name']} · {data['timestamp'] or 'N/A'}\nGenerated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"

        # Summary Slide
        if "summary" in include_sections:
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            title.text = "📊 Summary"
            
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.text = f"Site: {data['site_name']}"
            
            p = tf.add_paragraph()
            p.text = f"Total Changes: {data['total_changes']}"
            p.level = 1
            
            p = tf.add_paragraph()
            p.text = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            p.level = 1

        # Changes Slide
        if "changes" in include_sections and data.get("changes"):
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            title.text = "📝 Detected Changes"
            
            body = slide.placeholders[1]
            tf = body.text_frame
            
            for i, change in enumerate(data["changes"][:10]):
                if i == 0:
                    tf.text = f"[{change.get('severity', 'N/A')}] {change.get('type', 'N/A')}"
                else:
                    p = tf.add_paragraph()
                    p.text = f"[{change.get('severity', 'N/A')}] {change.get('type', 'N/A')}"
                
                p = tf.add_paragraph()
                p.text = f"URL: {change.get('url', 'N/A')[:60]}..."
                p.level = 1
                
                p = tf.add_paragraph()
                p.text = f"Before: {str(change.get('before', ''))[:50]}..."
                p.level = 2
                
                p = tf.add_paragraph()
                p.text = f"After: {str(change.get('after', ''))[:50]}..."
                p.level = 2
                
                if i < len(data["changes"]) - 1 and i < 9:
                    p = tf.add_paragraph()
                    p.text = ""

        # Insights Slide
        if "insights" in include_sections and data.get("povs"):
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            title.text = "💡 Insights & POVs"
            
            body = slide.placeholders[1]
            tf = body.text_frame
            
            for i, pov in enumerate(data["povs"][:5]):
                p = tf.add_paragraph()
                p.text = f"[{pov.get('priority', 'N/A').upper()}] {pov.get('observation', '')[:80]}..."
                
                p = tf.add_paragraph()
                p.text = f"Action: {pov.get('action', '')[:70]}..."
                p.level = 1
                
                if i < len(data["povs"]) - 1 and i < 4:
                    p = tf.add_paragraph()
                    p.text = ""

        # Actions Slide
        if "actions" in include_sections:
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            title.text = "🎯 Recommended Actions"
            
            body = slide.placeholders[1]
            tf = body.text_frame
            
            actions = [
                "Monitor critical changes daily",
                "Review SEO/AEO optimization opportunities",
                "Compare with Samsung.com/sg strategy",
                "Schedule follow-up analysis",
            ]
            
            for i, action in enumerate(actions):
                if i == 0:
                    tf.text = f"{i+1}. {action}"
                else:
                    p = tf.add_paragraph()
                    p.text = f"{i+1}. {action}"

        # Save presentation
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"apple_tracker_{data['site_name']}_{timestamp}.pptx"
        filepath = self.output_dir / filename
        
        prs.save(filepath)
        
        return str(filepath)

    async def export_to_png(
        self,
        data: Dict[str, Any],
        include_sections: List[str] = None,
    ) -> str:
        """
        Export dashboard data to PNG image.
        Note: This creates a simple text-based image.
        For full dashboard screenshot, use browser automation.
        
        Args:
            data: Dashboard export data
            include_sections: ['summary', 'changes', 'insights', 'actions']
            
        Returns:
            Path to generated PNG file
        """
        # For PNG export, we'll create a simple summary image
        # For full dashboard screenshots, use the browser-based approach
        from PIL import Image, ImageDraw, ImageFont
        
        # Create image
        width = 800
        height = 600
        image = Image.new('RGB', (width, height), color='white')
        draw = ImageDraw.Draw(image)
        
        # Try to load a font
        try:
            font_title = ImageFont.truetype("arial.ttf", 24)
            font_text = ImageFont.truetype("arial.ttf", 14)
        except:
            font_title = ImageFont.load_default()
            font_text = ImageFont.load_default()
        
        # Draw title
        draw.text((20, 20), f"🍎 Apple Tracker Report", fill='black', font=font_title)
        draw.text((20, 50), f"Site: {data['site_name']}", fill='gray', font=font_text)
        draw.text((20, 70), f"Total Changes: {data['total_changes']}", fill='gray', font=font_text)
        
        # Draw changes summary
        y = 110
        draw.text((20, y), "Recent Changes:", fill='black', font=font_text)
        y += 25
        
        for change in data.get("changes", [])[:5]:
            severity = change.get('severity', 'N/A')
            change_type = change.get('type', 'N/A')
            url = change.get('url', 'N/A')[:50]
            
            color = 'red' if severity == 'Critical' else 'orange' if severity == 'High' else 'green'
            draw.text((30, y), f"• [{severity}] {change_type}", fill=color, font=font_text)
            y += 18
        
        # Draw footer
        y = height - 40
        draw.text((20, y), f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", fill='gray', font=font_text)
        
        # Save image
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"apple_tracker_{data['site_name']}_{timestamp}.png"
        filepath = self.output_dir / filename
        
        image.save(filepath)
        
        return str(filepath)


# Global export service instance
export_service = DashboardExportService()


def get_export_service() -> DashboardExportService:
    """Get the export service instance"""
    return export_service
