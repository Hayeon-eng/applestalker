"""
spec_rule_db.py — 큐비 — Spec QA Rule DB [V2]

Rule 기반 Spec Validation 엔진의 룰 저장소.
엑셀(Rule DB 시트들) → 표준 JSON 룰셋으로 파싱하고, DB(qb_spec_rules)에
제품별 1행(JSON 문서)으로 저장한다. DB가 비어 있으면 저장소의 시드 JSON을 쓴다.

시트 → JSON 매핑
  · *MasterSpec*   → rules[]        (Rule ID/Category/Attribute/Official Value/Unit/
                                      Validation/Priority/Page/Interaction/Exception/Fix Guide)
  · Dictionary     → dictionary{}   (Representative → [aliases])
  · ExceptionRule  → exceptions[]   (rule/description)
  · InteractionRule→ interactions[] (section/action)  — Playwright 단계에서 사용
  · CountryException(선택) → country_exceptions[]
  · Candidates 는 런타임 승인 대기 목록(candidates[])로 관리

설계 원칙: 코드가 아니라 이 데이터만 갱신하면 룰이 확장된다.
"""
from __future__ import annotations
import io
import json
import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional

_HERE = os.path.dirname(__file__)

# 제품 목록에서 항상 제외되는 가짜/공통 키.
# [2026-07 FIX] Global Dictionary 시드가 spec_rules.seed.global.json 이름으로 저장돼
# 제품 필터에 'global'이라는 유령 제품 칩이 떴다 — 시드 파일명을 __global__로 통일하고
# (spec_dict_global.load()가 찾는 이름과 일치), 목록에서는 두 표기를 모두 제외한다.
_NON_PRODUCT_KEYS = {"__global__", "global", "__deleted__"}

# Validation 타입 표준화(엑셀 표기 흔들림 흡수)
_VALIDATION_TYPES = {
    "exact": "exact",
    "numeric exact": "numeric_exact",
    "numeric": "numeric_exact",
    "prefix match": "prefix",
    "prefix": "prefix",
    "dictionary": "dictionary",
    "option match": "option_match",
    "options": "option_match",
    "exists": "exists",
    "exist": "exists",
}


def _norm_header(h) -> str:
    return re.sub(r"[^a-z]", "", str(h or "").lower())


def _cell(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("-", "—") else s


def parse_xlsx(content: bytes, product: str, version: str = "") -> Dict[str, Any]:
    """Rule DB 엑셀(bytes) → 표준 룰셋 JSON.
    시트명은 부분 일치로 찾는다(MasterSpec/Dictionary/Exception/Interaction/Country)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)

    def _find_sheet(*keywords):
        for name in wb.sheetnames:
            low = name.lower()
            if any(k in low for k in keywords):
                return wb[name]
        return None

    ruleset: Dict[str, Any] = {
        "product": product, "version": version or "uploaded",
        "rules": [], "dictionary": {}, "exceptions": [],
        "interactions": [], "country_exceptions": [], "candidates": [],
        # [V3] 단위 동의어(운영자 관리 시트) · 전작 정답지(전작 비교 문구 판정용)
        "unit_synonyms": {}, "previous_models": [],
    }

    # ── MasterSpec ──
    ws = _find_sheet("masterspec", "master")
    if ws is None:
        raise ValueError("MasterSpec 시트를 찾을 수 없습니다.")
    rows = ws.iter_rows(values_only=True)
    header = [_norm_header(h) for h in next(rows)]

    def col(row, *names):
        for n in names:
            if n in header:
                return _cell(row[header.index(n)])
        return ""

    for row in rows:
        if not row or not _cell(row[0]):
            continue
        vtype_raw = col(row, "validation").lower()
        vtype = _VALIDATION_TYPES.get(vtype_raw, vtype_raw or "exact")
        # [V3 운영 결정] exists(Present/Included류) 룰 전면 삭제 — "값 없음 ≠ 오류"
        # 원칙에 따라 존재 확인형 룰은 기준표에서 제외한다. 엑셀에 남아 있어도 무시.
        if vtype == "exists":
            continue
        rule = {
            "rule_id": col(row, "ruleid"),
            "category": col(row, "category"),
            "attribute": col(row, "attribute"),
            # [V3] Official Value는 '|'로 복수 정답 허용 — 예: '4400|4272', '8.0|7.8'
            "expected": col(row, "officialvaluedraft", "officialvalue", "value"),
            "unit": col(row, "unit"),
            "validation": vtype,
            # [V3] Qualifier 컬럼 — '4400=typical,일반,標準;4272=rated,정격,定格'
            # 값별 한정어 오짝(숫자는 집합 안인데 라벨이 다른 값의 것) FAIL 판정용.
            "qualifier": col(row, "qualifier", "qualifiers"),
            "priority": col(row, "priority") or "Medium",
            "page": col(row, "page") or "PDP",
            "interaction": col(row, "interaction"),
            "exception": col(row, "exception"),
            "fix_guide": col(row, "fixguide"),
            "notes": col(row, "notes"),
        }
        if rule["rule_id"] and rule["attribute"]:
            ruleset["rules"].append(rule)

    # ── Dictionary ──
    ws = _find_sheet("dictionary")
    if ws is not None:
        rows = ws.iter_rows(values_only=True)
        next(rows, None)  # header
        for row in rows:
            if not row:
                continue
            rep, alias = _cell(row[0]), _cell(row[1] if len(row) > 1 else "")
            if rep and alias:
                ruleset["dictionary"].setdefault(rep, [])
                if alias not in ruleset["dictionary"][rep]:
                    ruleset["dictionary"][rep].append(alias)

    # ── ExceptionRule ──
    ws = _find_sheet("exception")
    if ws is not None and "country" not in (ws.title or "").lower():
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not row or not _cell(row[0]):
                continue
            ruleset["exceptions"].append({
                "rule": _cell(row[0]),
                "description": _cell(row[1] if len(row) > 1 else ""),
            })

    # ── InteractionRule ──
    ws = _find_sheet("interaction")
    if ws is not None:
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not row or not _cell(row[0]):
                continue
            ruleset["interactions"].append({
                "section": _cell(row[0]),
                "action": _cell(row[1] if len(row) > 1 else ""),
            })

    # ── CountryException(선택) ──
    ws = _find_sheet("countryexception", "country")
    if ws is not None:
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not row or not _cell(row[0]):
                continue
            ruleset["country_exceptions"].append({
                "country": _cell(row[0]),
                "rule": _cell(row[1] if len(row) > 1 else ""),
                "action": _cell(row[2] if len(row) > 2 else ""),
            })

    # ── UnitSynonym(선택) — [V3] 운영자 관리 다국어 단위 표기 시트 ──
    #    컬럼: Canonical Unit | Synonym  (예: nits | 니트 · nits | ニト · inch | Zoll)
    #    코드 기본 테이블(spec_value_match.DEFAULT_UNIT_SYNONYMS)에 병합·우선된다.
    ws = _find_sheet("unitsynonym", "unit synonym", "단위")
    if ws is not None:
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if not row:
                continue
            canon, syn = _cell(row[0]), _cell(row[1] if len(row) > 1 else "")
            if canon and syn:
                ruleset["unit_synonyms"].setdefault(canon, [])
                if syn not in ruleset["unit_synonyms"][canon]:
                    ruleset["unit_synonyms"][canon].append(syn)

    # ── PreviousModel(선택) — [V3] 전작 정답지 시트 ──
    #    컬럼: Model | Aliases(콤마) | Attribute | Values('|' 복수) | Unit
    #    PDP의 전작 비교 문구("Fold6의 1,750니트보다…") 속 숫자를 전작 정답과
    #    대조해 pass/fail 판정하기 위한 데이터. 직전 1세대만 운영(운영 결정).
    ws = _find_sheet("previousmodel", "previous model", "전작")
    if ws is not None:
        rows = ws.iter_rows(values_only=True)
        next(rows, None)
        by_model: Dict[str, Dict[str, Any]] = {}
        for row in rows:
            if not row or not _cell(row[0]):
                continue
            model = _cell(row[0])
            aliases = [a.strip() for a in _cell(row[1] if len(row) > 1 else "").split(",") if a.strip()]
            attribute = _cell(row[2] if len(row) > 2 else "")
            values = [v.strip() for v in _cell(row[3] if len(row) > 3 else "").split("|") if v.strip()]
            unit = _cell(row[4] if len(row) > 4 else "")
            pm = by_model.setdefault(model, {"model": model, "aliases": aliases, "specs": []})
            for a in aliases:
                if a not in pm["aliases"]:
                    pm["aliases"].append(a)
            if attribute and values:
                pm["specs"].append({"attribute": attribute, "values": values, "unit": unit})
        ruleset["previous_models"] = list(by_model.values())

    return ruleset


# ── DB 저장/로드 (qb_history와 동일한 SessionLocal 재사용) ──────────────
def _db():
    import sys
    sys.path.insert(0, os.path.join(_HERE, ".."))
    from database import SessionLocal
    from models import QbSpecRules
    return SessionLocal, QbSpecRules


def save(product: str, ruleset: Dict[str, Any], version: str = "", merge_dictionary: bool = True) -> Dict[str, Any]:
    """[2026-07 버그 수정] 기존엔 재업로드 시 dictionary를 통째로 덮어써서, MasterSpec(스펙
    값)만 고치려고 엑셀을 재업로드해도 그동안 운영자가 승인해둔 alias가 전부 삭제됐다.
    merge_dictionary=True(기본)면 저장 전에 '기존 DB의 dictionary ∪ 새로 넘어온 dictionary'로
    합쳐서 저장한다 — 새 alias는 추가되고, 예전에 승인된 alias는 절대 사라지지 않는다.
    (Global Dictionary 자체를 저장할 때는 spec_dict_global.save()가 merge_dictionary=False로
    호출한다 — 이미 병합이 끝난 최종본이므로 이중 병합을 피하기 위함.)"""
    SessionLocal, QbSpecRules = _db()
    db = SessionLocal()
    try:
        row = db.query(QbSpecRules).filter(QbSpecRules.product == product).first()
        if row and merge_dictionary and row.data:
            try:
                prev = json.loads(row.data)
                merged = dict(ruleset.get("dictionary", {}))
                for rep, aliases in (prev.get("dictionary") or {}).items():
                    cur = merged.setdefault(rep, [])
                    for a in aliases:
                        if a not in cur:
                            cur.append(a)
                ruleset = {**ruleset, "dictionary": merged}
            except Exception as e:
                print(f"[spec_rule_db] dictionary merge skip (falling back to overwrite): {e}")
        elif not row and merge_dictionary:
            # [2026-07 FIX] 제품 삭제 후 재업로드 — delete_product()가 보관해둔 Dictionary를
            # 자동으로 복원 병합한다(삭제 이전에 승인해둔 alias가 유실되지 않도록).
            backup = _load_dictionary_backup(product)
            if backup:
                merged = dict(ruleset.get("dictionary", {}))
                for rep, aliases in backup.items():
                    cur = merged.setdefault(rep, [])
                    for a in aliases:
                        if a not in cur:
                            cur.append(a)
                ruleset = {**ruleset, "dictionary": merged}
        payload = json.dumps(ruleset, ensure_ascii=False)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if row:
            row.data = payload
            row.version = version or ruleset.get("version", "")
            row.updated_at = now
        else:
            row = QbSpecRules(product=product, data=payload,
                              version=version or ruleset.get("version", ""), updated_at=now)
            db.add(row)
        db.commit()
        # [2026-07] 삭제 마커에 있던 제품을 다시 업로드하면 마커에서 자동 해제
        try:
            dp = deleted_products()
            if product in dp:
                _set_deleted([p for p in dp if p != product])
        except Exception:
            pass
        return {"product": product, "version": row.version, "updated_at": now,
                "rules": len(ruleset.get("rules", []))}
    finally:
        db.close()


def load(product: str) -> Optional[Dict[str, Any]]:
    """DB → 없으면 저장소 시드(spec_rules.seed.{product}.json) → 없으면 None."""
    try:
        SessionLocal, QbSpecRules = _db()
        db = SessionLocal()
        try:
            row = db.query(QbSpecRules).filter(QbSpecRules.product == product).first()
            if row and row.data:
                return json.loads(row.data)
        finally:
            db.close()
    except Exception as e:  # DB 미설정(로컬 등)이어도 시드로 동작
        print(f"[spec_rule_db] DB load skip: {e}")
    if product in deleted_products():  # 삭제된 제품 — 시드 폴백 금지
        return None
    seed = os.path.join(_HERE, f"spec_rules.seed.{product}.json")
    if os.path.exists(seed):
        return json.load(open(seed, encoding="utf-8"))
    return None


def list_products() -> List[Dict[str, Any]]:
    out, seen = [], set()
    deleted = set(deleted_products())
    try:
        SessionLocal, QbSpecRules = _db()
        db = SessionLocal()
        try:
            for row in db.query(QbSpecRules).all():
                if row.product in _NON_PRODUCT_KEYS or row.product.startswith("__dict_backup__"):
                    continue  # Global Dictionary/마커/용어사전 백업 행 — 목록 제외
                out.append({"product": row.product, "version": row.version,
                            "updated_at": row.updated_at, "source": "db"})
                seen.add(row.product)
        finally:
            db.close()
    except Exception:
        pass
    for fn in os.listdir(_HERE):
        m = re.match(r"spec_rules\.seed\.(.+)\.json$", fn)
        if not m:
            continue
        p = m.group(1)
        # [2026-07] 삭제된 제품(deleted_products)은 시드가 저장소에 남아 있어도 되살리지 않는다.
        if p in seen or p in _NON_PRODUCT_KEYS or p in deleted:
            continue
        out.append({"product": p, "version": "seed", "updated_at": "", "source": "seed"})
    return out


# ── [2026-07 신규] 제품 Rule DB 삭제 ────────────────────────────────────
# 신모델(예: 다음 세대 폴더블)을 엑셀로 올린 뒤 구모델(Fold7/Flip7 등)을 목록에서
# 내리기 위한 기능. DB 행을 지우고, 저장소에 시드 JSON이 남아 있어도 부활하지 않도록
# '__deleted__' 마커 행(제품명 목록 JSON)에 기록한다. 같은 제품을 다시 엑셀 업로드하면
# 마커에서 자동 해제된다(save 참조).

def deleted_products() -> List[str]:
    try:
        SessionLocal, QbSpecRules = _db()
        db = SessionLocal()
        try:
            row = db.query(QbSpecRules).filter(QbSpecRules.product == "__deleted__").first()
            if row and row.data:
                return list(json.loads(row.data))
        finally:
            db.close()
    except Exception:
        pass
    return []


def _set_deleted(products: List[str]) -> None:
    SessionLocal, QbSpecRules = _db()
    db = SessionLocal()
    try:
        row = db.query(QbSpecRules).filter(QbSpecRules.product == "__deleted__").first()
        payload = json.dumps(sorted(set(products)), ensure_ascii=False)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if row:
            row.data = payload; row.updated_at = now
        else:
            db.add(QbSpecRules(product="__deleted__", data=payload, version="marker", updated_at=now))
        db.commit()
    finally:
        db.close()


def _dict_backup_key(product: str) -> str:
    return f"__dict_backup__{product}"


def _backup_dictionary(product: str, dictionary: Dict[str, List[str]]) -> None:
    """제품 삭제 전, 그 제품의 Dictionary만 별도 보관 행에 병합 저장.
    save()가 재업로드 시 이 보관본과 자동 병합해 복원한다(add_alias 등으로 쌓인
    승인 이력이 제품 삭제로 함께 사라지지 않도록)."""
    if not dictionary:
        return
    SessionLocal, QbSpecRules = _db()
    db = SessionLocal()
    try:
        key = _dict_backup_key(product)
        row = db.query(QbSpecRules).filter(QbSpecRules.product == key).first()
        merged = dict(dictionary)
        if row and row.data:
            try:
                prev = json.loads(row.data) or {}
                for rep, aliases in prev.items():
                    cur = merged.setdefault(rep, [])
                    for a in aliases:
                        if a not in cur:
                            cur.append(a)
            except Exception:
                pass
        payload = json.dumps(merged, ensure_ascii=False)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        if row:
            row.data = payload; row.updated_at = now
        else:
            db.add(QbSpecRules(product=key, data=payload, version="dict_backup", updated_at=now))
        db.commit()
    finally:
        db.close()


def _load_dictionary_backup(product: str) -> Dict[str, List[str]]:
    try:
        SessionLocal, QbSpecRules = _db()
        db = SessionLocal()
        try:
            row = db.query(QbSpecRules).filter(QbSpecRules.product == _dict_backup_key(product)).first()
            if row and row.data:
                return json.loads(row.data)
        finally:
            db.close()
    except Exception:
        pass
    return {}


def delete_product(product: str) -> Dict[str, Any]:
    """제품의 Rule DB 삭제(룰·예외·후보 등 제품 데이터만). [2026-07 FIX] 제품 용어사전
    (Dictionary)은 여기서 함께 지우지 않는다 — 삭제 전 별도 보관 행으로 백업해두었다가,
    같은 제품을 다시 업로드하면 save()가 자동으로 복원한다. Dictionary 자체를 완전히
    지우려면 delete_dictionary()를 별도로 호출해야 한다. Global(공통) Dictionary와
    다른 제품에는 영향이 없다."""
    if product in _NON_PRODUCT_KEYS:
        raise ValueError("공통/마커 행은 삭제할 수 없습니다.")
    removed_db = False
    try:
        SessionLocal, QbSpecRules = _db()
        db = SessionLocal()
        try:
            row = db.query(QbSpecRules).filter(QbSpecRules.product == product).first()
            if row and row.data:
                try:
                    prev = json.loads(row.data)
                    _backup_dictionary(product, prev.get("dictionary") or {})
                except Exception as e:
                    print(f"[spec_rule_db] dictionary backup skip: {e}")
            n = db.query(QbSpecRules).filter(QbSpecRules.product == product).delete(synchronize_session=False)
            db.commit()
            removed_db = n > 0
        finally:
            db.close()
    except Exception as e:
        print(f"[spec_rule_db] delete skip: {e}")
    has_seed = os.path.exists(os.path.join(_HERE, f"spec_rules.seed.{product}.json"))
    if has_seed:  # 시드 부활 방지 마커
        _set_deleted(deleted_products() + [product])
    return {"product": product, "removed_db": removed_db, "seed_blocked": has_seed}


def delete_dictionary(product: str) -> Dict[str, Any]:
    """[2026-07 신규] 제품 용어사전(Dictionary)만 별도로 완전 삭제 — 제품 삭제와는 분리된
    기능이다. 보관본(백업 행)과, 아직 남아있는 제품 룰 행 안의 dictionary를 모두 비운다.
    Global(공통) Dictionary와 다른 제품에는 영향이 없다."""
    SessionLocal, QbSpecRules = _db()
    db = SessionLocal()
    removed = False
    try:
        n = db.query(QbSpecRules).filter(QbSpecRules.product == _dict_backup_key(product)).delete(synchronize_session=False)
        removed = n > 0
        row = db.query(QbSpecRules).filter(QbSpecRules.product == product).first()
        if row and row.data:
            try:
                data = json.loads(row.data)
                if data.get("dictionary"):
                    data["dictionary"] = {}
                    row.data = json.dumps(data, ensure_ascii=False)
                    row.updated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    removed = True
            except Exception as e:
                print(f"[spec_rule_db] delete_dictionary row clear skip: {e}")
        db.commit()
    finally:
        db.close()
    return {"product": product, "dictionary_removed": removed}


def add_alias(product: str, representative: str, alias: str) -> Dict[str, Any]:
    """Dictionary에 Alias 추가(사용자 승인 후 호출). Candidate에 있으면 제거."""
    rs = load(product)
    if not rs:
        raise ValueError(f"룰셋 없음: {product}")
    rs["dictionary"].setdefault(representative, [])
    if alias not in rs["dictionary"][representative]:
        rs["dictionary"][representative].append(alias)
    rs["candidates"] = [c for c in rs.get("candidates", [])
                        if not (c.get("alias") == alias and c.get("representative") == representative)]
    save(product, rs, version=rs.get("version", ""))
    return {"representative": representative, "aliases": rs["dictionary"][representative]}
