"""
copy_rule_parser.py — 큐비 — Dotcom QA 체커 [Phase C]

큐비 🐝 — 풀네임 QA Bee, 줄여서 큐비. 닷컴 페이지를 붕붕 돌며 규칙대로 검수.

Samsung 카피덱(copydeck.xlsx: M3/M12 시트)에서 '언어 불변' 검수 토큰을 뽑아
copy_rules.json 으로 만든다. 번역 사이트에서도 안 바뀌는 값만 정확 검사 대상으로 삼는다.

- spec_tokens : 하드웨어 단위가 붙은 숫자 (200 MP, 1 TB, 512 GB, 5000 mAh, 2600 nits, 100x ...)
                → 번역돼도 동일하므로 페이지에 '그대로 있는지' 정확 검사
- proper_nouns: 번역돼도 영문 유지되는 고유명사 후보 (Snapdragon 8 Elite Gen 5, Vapor Chamber ...)
                → 자동 추출은 후보 수준이라 '사람이 검토'하는 하이브리드

서술형 문장은 언어 가변이라 여기서 다루지 않는다(체커에서 존재/placeholder 여부만 별도로).
"""
from __future__ import annotations
import json
import re
from typing import Any, Dict, List

# 하드웨어 단위(타이트하게). 숫자와 단위 사이 공백은 선택. 콤마는 허용하지 않아 날짜(2026,)·천단위 노이즈 배제.
_UNIT = r"(?:MP|GB|TB|mAh|Wh|Hz|nits?|inch|W|mm)"
_SPEC_RE = re.compile(r"(?<![\w.])(\d+(?:\.\d+)?)\s?(" + _UNIT + r')(?![\w])', re.IGNORECASE)
# 줌 배율(100x, 10x, 2x ...) — 카메라 스펙
_ZOOM_RE = re.compile(r"(?<![\w.])(\d+(?:\.\d+)?)\s?([x×])(?![\w])")

# 고유명사 후보: 브랜드/기술 키워드로 시작하거나 포함하는 TitleCase 구
_PN_SEED = re.compile(
    r"\b("
    r"Snapdragon\s8\sElite(?:\sGen\s?\d+)?|"
    r"Exynos\s?\d{3,4}|"
    r"Galaxy\s(?:S26\sUltra|S26\+|S26|Buds\d?\sPro|Watch|AI)|Galaxy\sAI|"
    r"Vapor\sChamber|Corning\sGorilla\s(?:Glass\sVictus|Armor)\s?\d?|"
    r"Nightography|ProVisual\sEngine|Privacy\sDisplay|Now\sBrief|Photo\sAssist"
    r")\b"
)

# 고유명사 후보에서 걸러낼 노이즈(문장 조각 등)
_PN_STOP = re.compile(r"\b(and|while|have|has|been|for|the|with|of)\b", re.IGNORECASE)


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", t.strip())


def extract_from_cells(cells: List[str]) -> Dict[str, List[str]]:
    specs, zooms, pns = set(), set(), set()
    for c in cells:
        for m in _SPEC_RE.finditer(c):
            num, unit = m.group(1), m.group(2)
            unit = "nits" if unit.lower().startswith("nit") else unit
            specs.add(_norm(f"{num} {unit}"))
        for m in _ZOOM_RE.finditer(c):
            zooms.add(_norm(f"{m.group(1)}{ 'x' }"))
        for m in _PN_SEED.finditer(c):
            pn = _norm(m.group(0))
            if 4 <= len(pn) <= 40 and not _PN_STOP.search(pn):
                pns.add(pn)
    spec_tokens = sorted(specs, key=lambda s: (s.split()[-1], float(re.findall(r"[\d.]+", s)[0]))) + sorted(zooms)
    return {"spec_tokens": spec_tokens, "proper_nouns": sorted(pns)}


def build_rules(xlsx_path: str, sheets=("M3", "M12")) -> Dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    out: Dict[str, Any] = {"source": xlsx_path.rsplit("/", 1)[-1], "products": {}}
    for sheet in sheets:
        if sheet not in wb.sheetnames:
            continue
        cells = [str(v) for r in wb[sheet].iter_rows(values_only=True) for v in r if v]
        out["products"][sheet] = extract_from_cells(cells)
    wb.close()
    return out


if __name__ == "__main__":
    import sys
    src = sys.argv[1] if len(sys.argv) > 1 else "/mnt/user-data/uploads/copydeck.xlsx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "copy_rules.json"
    rules = build_rules(src)
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(rules, f, ensure_ascii=False, indent=2)
    for prod, d in rules["products"].items():
        print(f"[{prod}] spec_tokens={len(d['spec_tokens'])} proper_nouns={len(d['proper_nouns'])}")
        print("   specs:", ", ".join(d["spec_tokens"]))
        print("   nouns:", ", ".join(d["proper_nouns"]))
    print(f"\n→ {dst} 저장 (proper_nouns 는 사람이 한 번 검토 권장)")
