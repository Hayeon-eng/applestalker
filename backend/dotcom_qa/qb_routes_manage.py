"""
qb_routes_manage.py — 사이트·제품·구식 룰 관리 (/sites*, /products*, /specs*, /rules/*/add) [qb_api 분할]
V2 제품(fold7/flip7 등)의 구식 편집은 여기서 409로 차단된다 — 원본은 Rule DB 하나.
"""
from __future__ import annotations
import asyncio
import json
import os
import sys
from datetime import datetime
from typing import Any, Callable, Dict, List, Optional

sys.path.insert(0, os.path.dirname(__file__))

from fastapi import APIRouter, Body, HTTPException, Query
from fastapi.responses import JSONResponse, StreamingResponse, HTMLResponse

import runner
import qa_report
import schema_checker
import html_qa_scoring
from database import SessionLocal
from models import QbHistory

import qb_core
from qb_core import qb_router, _registry, PAGE_TYPES, SCHEMA_TYPES

@qb_router.get("/sites")
def qb_sites():
    by = _registry.by_region_unique()  # 사이트코드 중복 제거 — 지역 칩엔 로케일 1개씩
    return {"count": len(_registry.unique_sitecodes()),   # 사이트(로케일) 수
            "page_count": len(_registry.all()),           # 실제 크롤 대상(제품×페이지타입) 수
            "regions": {r: [{"sitecode": s["sitecode"], "country": s["country"],
                             "lang": s["lang"], "url": s["url"]} for s in sites]
                        for r, sites in by.items()},
            "missing_lang": _registry.missing_lang()}


# PAGE_TYPES / SCHEMA_TYPES 는 qb_core에서 import (이 모듈의 SPEC_CATALOG 등과 함께 사용)


# ── 표준 스펙 항목 사전(드롭다운 + 예시) ──
SPEC_CATALOG = [
    {"category": "display_size", "label": "디스플레이 크기", "ex_value": "6.9", "ex_unit": "inch"},
    {"category": "brightness", "label": "밝기", "ex_value": "2600", "ex_unit": "nits"},
    {"category": "refresh_rate", "label": "주사율", "ex_value": "120", "ex_unit": "Hz"},
    {"category": "camera_wide", "label": "후면 메인(Wide)", "ex_value": "200", "ex_unit": "MP"},
    {"category": "camera_ultrawide", "label": "울트라와이드", "ex_value": "50", "ex_unit": "MP"},
    {"category": "camera_telephoto", "label": "망원", "ex_value": "50", "ex_unit": "MP"},
    {"category": "camera_front", "label": "전면 카메라", "ex_value": "12", "ex_unit": "MP"},
    {"category": "space_zoom", "label": "공간 줌", "ex_value": "100", "ex_unit": "x"},
    {"category": "video_playback", "label": "영상 재생(배터리)", "ex_value": "31", "ex_unit": "hours"},
    {"category": "battery_capacity", "label": "배터리 용량", "ex_value": "5000", "ex_unit": "mAh"},
    {"category": "charging", "label": "고속충전", "ex_value": "45", "ex_unit": "W"},
    {"category": "memory", "label": "메모리", "ex_value": "16", "ex_unit": "GB"},
    {"category": "storage", "label": "저장", "ex_value": "512", "ex_unit": "GB"},
    {"category": "processor", "label": "프로세서", "ex_value": "Snapdragon 8 Elite Gen 5", "ex_unit": ""},
    {"category": "display_type", "label": "디스플레이 종류", "ex_value": "Dynamic AMOLED 2X", "ex_unit": ""},
    {"category": "camera_periscope", "label": "페리스코프 망원", "ex_value": "50", "ex_unit": "MP"},
    {"category": "wireless_charging", "label": "무선충전", "ex_value": "25", "ex_unit": "W"},
    # ── 버즈(Buds) 계열 ──
    {"category": "battery_buds", "label": "이어버드 배터리", "ex_value": "61", "ex_unit": "mAh"},
    {"category": "battery_case", "label": "케이스 배터리", "ex_value": "530", "ex_unit": "mAh"},
    {"category": "playback_anc_on", "label": "재생(ANC 켬)", "ex_value": "6", "ex_unit": "hours"},
    {"category": "playback_anc_off", "label": "재생(ANC 끔)", "ex_value": "7", "ex_unit": "hours"},
    {"category": "playback_with_case", "label": "재생(케이스 포함)", "ex_value": "30", "ex_unit": "hours"},
    {"category": "audio_bit", "label": "오디오 비트", "ex_value": "24", "ex_unit": "bit"},
    {"category": "audio_khz", "label": "오디오 샘플레이트", "ex_value": "96", "ex_unit": "kHz"},
    {"category": "ip_rating", "label": "방수방진 등급", "ex_value": "IP57", "ex_unit": ""},
]


@qb_router.get("/spec-catalog")
def qb_spec_catalog():
    return {"catalog": SPEC_CATALOG}


@qb_router.get("/products")
def qb_products():
    data = _load_specs().get("products", {})
    out = []
    for code, entry in data.items():
        label = entry.get("label", code) if isinstance(entry, dict) else code
        spec_only = entry.get("spec_only", False) if isinstance(entry, dict) else False
        out.append({"code": code, "label": label, "spec_only": spec_only})
    # [V2] Rule DB(spec_rules) 제품도 드롭다운에 합류 — 없으면 V2 화면으로 갈 입구가 없음
    try:
        import spec_rule_db as _srd

        def _v2_label(code: str) -> str:
            # 슬러그 → 사람이 읽는 제품명. 새 계열이 나오면 여기 한 줄이면 된다.
            KNOWN = {
                "galaxy-watch8": "Galaxy Watch8",
                "galaxy-watch-ultra": "Galaxy Watch Ultra",
            }
            if code in KNOWN:
                return KNOWN[code]
            out = code.replace("galaxy-z-", "Galaxy Z ").replace("galaxy-watch", "Galaxy Watch")
            out = out.replace("fold", "Fold").replace("flip", "Flip")
            return out
        have = {p["code"] for p in out}
        for p in _srd.list_products():
            if p["product"] not in have:
                out.append({"code": p["product"], "label": _v2_label(p["product"]), "spec_only": False, "v2": True})
            else:
                for o in out:
                    if o["code"] == p["product"]:
                        o["v2"] = True
    except Exception as e:
        print(f"[qb_products] v2 merge skip: {e}")
    # key_specs.json 의 등록 순서를 그대로 유지(정렬하지 않음)
    return {"products": out}


@qb_router.post("/products/add")
def qb_products_add(payload: Dict[str, Any] = Body(...)):
    name = (payload.get("product") or "").strip()
    if not name:
        raise HTTPException(400, "product가 필요합니다.")
    if qb_core._v2_managed(name):
        raise HTTPException(409, f"'{name}'은(는) 이미 Rule DB(V2) 제품으로 등록돼 있습니다. 신규 V2 제품은 Rule DB 엑셀 업로드로 추가하세요.")
    label = (payload.get("label") or name).strip()
    data = _load_specs()
    prods = data.setdefault("products", {})
    if name not in prods:
        prods[name] = {"label": label, "specs": []}
    _save_specs(data)
    return {"ok": True, "products": [{"code": c, "label": (e.get("label", c) if isinstance(e, dict) else c)} for c, e in prods.items()]}




@qb_router.post("/sites/add")
def qb_sites_add(payload: Dict[str, Any] = Body(...)):
    url = (payload.get("url") or "").strip()
    if not url:
        raise HTTPException(400, "url이 필요합니다.")
    sc = payload.get("sitecode") or qb_core._sitecode_from_url(url)
    if _registry.get(sc):
        _registry.update(sc, url=url)
    else:
        _registry.add(sc, url, region=payload.get("region", ""), country=payload.get("country", ""),
                      lang=payload.get("lang", ""), product=payload.get("product", "galaxy-s26-ultra"))
    _registry.save()
    return {"ok": True, "sitecode": sc}


@qb_router.post("/sites/remove")
def qb_sites_remove(payload: Dict[str, Any] = Body(...)):
    sc = payload.get("sitecode")
    ok = _registry.remove(sc)
    if ok:
        _registry.save()
    return {"ok": ok}


# ── 핵심 스펙(3축: product/category/value/unit) CRUD ──
_SPEC_PATH = os.path.join(os.path.dirname(__file__), "key_specs.json")


def _load_specs() -> Dict[str, Any]:
    try:
        return json.load(open(_SPEC_PATH, encoding="utf-8"))
    except Exception:
        return {"products": {}}


def _save_specs(data: Dict[str, Any]):
    with open(_SPEC_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _prod_entry(data, product):
    e = data.setdefault("products", {}).setdefault(product, {"label": product, "specs": []})
    if isinstance(e, list):  # 구버전 호환: 리스트면 감싸기
        e = {"label": product, "specs": e}; data["products"][product] = e
    e.setdefault("specs", [])
    return e


@qb_router.get("/specs")
def qb_specs(product: str = Query("galaxy-s26-ultra")):
    data = _load_specs().get("products", {})
    e = data.get(product, {})
    if isinstance(e, list):
        return {"product": product, "label": product, "specs": e}
    return {"product": product, "label": e.get("label", product), "specs": e.get("specs", [])}


def _v2_managed(product: str) -> bool:
    """이 제품의 기준값이 Rule DB(V2)로 관리되는지. V2 제품은 구식 편집 경로
    (key_specs/copy_rules 직접 수정)를 전부 차단해 원본을 하나로 유지한다."""
    try:
        import spec_rule_db as _srd
        return product in {p["product"] for p in _srd.list_products()}
    except Exception:
        return False


_V2_EDIT_BLOCKED = ("이 제품의 기준값은 Rule DB(V2)로 관리됩니다 — 화면 편집 불가. "
                    "값 수정은 스펙 탭의 'Rule DB 엑셀 업로드', 표현 추가는 'Dictionary 추가' 승인을 사용하세요.")


@qb_router.post("/specs/add")
def qb_specs_add(payload: Dict[str, Any] = Body(...)):
    product = payload.get("product", "galaxy-s26-ultra")
    if qb_core._v2_managed(product):
        raise HTTPException(409, qb_core._V2_EDIT_BLOCKED)
    # values: 여러 값(국별 variation) 허용 — 콤마/리스트 모두 수용
    raw = payload.get("values", payload.get("value", ""))
    if isinstance(raw, list):
        values = [str(v).strip() for v in raw if str(v).strip()]
    else:
        values = [v.strip() for v in str(raw).split(",") if v.strip()]
    pts = payload.get("page_types") or ["PDP"]
    if isinstance(pts, str):
        pts = [p.strip() for p in pts.split(",") if p.strip()] or ["PDP"]
    row = {"category": payload.get("category", ""), "values": values,
           "unit": payload.get("unit", ""), "page_types": pts}
    if not row["category"] or not values:
        raise HTTPException(400, "category와 value(값)가 필요합니다.")
    data = _load_specs()
    _prod_entry(data, product)["specs"].append(row)
    _save_specs(data)
    return {"ok": True, "specs": data["products"][product]["specs"]}


@qb_router.post("/specs/remove")
def qb_specs_remove(payload: Dict[str, Any] = Body(...)):
    product = payload.get("product", "galaxy-s26-ultra")
    if qb_core._v2_managed(product):
        raise HTTPException(409, qb_core._V2_EDIT_BLOCKED)
    idx = payload.get("index")
    data = _load_specs()
    e = data.get("products", {}).get(product, {})
    arr = e.get("specs", []) if isinstance(e, dict) else e
    if isinstance(idx, int) and 0 <= idx < len(arr):
        arr.pop(idx); _save_specs(data)
        return {"ok": True, "specs": arr}
    return {"ok": False, "specs": arr}


# ── 검수 룰 추가 (카피/스키마 별도) ──
_COPY_PATH = os.path.join(os.path.dirname(__file__), "copy_rules.json")
_SCHEMA_PATH = os.path.join(os.path.dirname(__file__), "schema_rules.json")


@qb_router.post("/rules/copy/add")
def qb_rules_copy_add(payload: Dict[str, Any] = Body(...)):
    """스펙 QA 룰 추가: spec_token 또는 proper_noun.
    [V2] Rule DB 관리 제품은 차단 — 기준값 원본은 Rule DB 하나여야 한다."""
    product = payload.get("product", "M3")
    if qb_core._v2_managed(product):
        raise HTTPException(409, qb_core._V2_EDIT_BLOCKED)
    kind = payload.get("kind", "spec")  # "spec" | "proper_noun"
    token = (payload.get("token") or "").strip()
    if not token:
        raise HTTPException(400, "token이 필요합니다.")
    data = json.load(open(_COPY_PATH, encoding="utf-8"))
    key = "spec_tokens" if kind == "spec" else "proper_nouns"
    arr = data["products"].setdefault(product, {}).setdefault(key, [])
    if token not in arr:
        arr.append(token)
    json.dump(data, open(_COPY_PATH, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    return {"ok": True, key: arr}


@qb_router.post("/rules/schema/add")
def qb_rules_schema_add(payload: Dict[str, Any] = Body(...)):
    """스키마 QA 룰 추가(페이지타입별): 선택 타입 블록에 속성 + 기대값·값종류 등록. 블록 없으면 생성.
    value_kind: url(가변 치환·정확) / enum(정확) / text(번역=확인만) / exists(존재만)"""
    product = payload.get("product", "M3")
    page_type = payload.get("page_type", "PDP")
    block_type = (payload.get("block_type") or payload.get("block") or "").strip()
    prop = (payload.get("property") or "").strip()
    id_slug = (payload.get("id_slug") or "").strip() or None
    value = (payload.get("value") or "").strip()
    kind = (payload.get("value_kind") or "exists").strip()   # url|enum|text|exists
    nested = (payload.get("nested") or "").strip() or None    # @id|@type|None
    if not block_type or not prop:
        raise HTTPException(400, "block_type과 property가 필요합니다.")
    path = os.path.join(os.path.dirname(__file__), f"schema_rules.{product}.{page_type}.json")
    if os.path.exists(path):
        data = json.load(open(path, encoding="utf-8"))
    else:
        data = {"source": "manual", "sitecode_token": "{SITECODE}", "page_type": page_type, "blocks": []}
    blocks = data.setdefault("blocks", [])
    hit = next((b for b in blocks if b.get("name") == block_type or block_type in (b.get("types") or [])), None)
    if not hit:
        hit = {"name": block_type, "types": [block_type], "id_pattern": None, "id_slug": id_slug,
               "required_properties": [], "optional_properties": [], "property_notes": {},
               "haspart_ids": [], "expected_values": {}, "conditional": None}
        blocks.append(hit)
    if prop not in hit.setdefault("required_properties", []):
        hit["required_properties"].append(prop)
    # 기대값·종류 저장 (exists 는 값 검사 안 함 → expected_values 에 안 넣음)
    if kind != "exists" and value:
        hit.setdefault("expected_values", {})[prop] = {"value": value, "kind": kind, "nested": nested}
    json.dump(data, open(path, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    return {"ok": True, "product": product, "page_type": page_type, "block": block_type,
            "required_properties": hit["required_properties"],
            "expected_values": hit.get("expected_values", {})}


# ── URL 엑셀 템플릿 다운로드 / 일괄 업로드 ──
_URL_COLS = ["sitecode", "url", "region", "country", "lang", "product"]


@qb_router.get("/sites/template.xlsx")
def qb_sites_template():
    from openpyxl import Workbook
    import io
    wb = Workbook(); ws = wb.active; ws.title = "URLs"
    ws.append(_URL_COLS)
    ws.append(["sg", "https://www.samsung.com/sg/smartphones/galaxy-s26-ultra/compare/",
               "APAC", "Singapore", "en-SG", "galaxy-s26-ultra"])  # 예시 1행
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": "attachment; filename=qubi_url_template.xlsx"})


@qb_router.post("/sites/upload")
def qb_sites_upload(payload: Dict[str, Any] = Body(...)):
    """base64 로 받은 xlsx(위 템플릿 형식)를 파싱해 일괄 추가."""
    import base64
    import io
    from openpyxl import load_workbook
    b64 = payload.get("b64") or ""
    if "," in b64:
        b64 = b64.split(",", 1)[1]  # dataURL 접두 제거
    try:
        wb = load_workbook(io.BytesIO(base64.b64decode(b64)), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(400, f"엑셀을 읽을 수 없습니다: {e}")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"ok": True, "added": 0}
    header = [str(c or "").strip().lower() for c in rows[0]]
    idx = {c: header.index(c) for c in _URL_COLS if c in header}
    added = 0
    for r in rows[1:]:
        def g(col):
            i = idx.get(col)
            return (str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else "")
        url = g("url")
        if not url:
            continue
        sc = g("sitecode") or qb_core._sitecode_from_url(url)
        if _registry.get(sc):
            _registry.update(sc, url=url)
        else:
            _registry.add(sc, url, region=g("region"), country=g("country"),
                          lang=g("lang"), product=g("product") or "galaxy-s26-ultra")
        added += 1
    _registry.save()
    return {"ok": True, "added": added, "count": len(_registry.all())}


