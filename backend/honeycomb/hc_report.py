"""hc_report.py — honeyComb 엑셀(기존 Google_Shopping_GMC_Final_Report 시트 구조 유지) [2026-09]"""
from __future__ import annotations
import io
from typing import Any, Dict, List

import hc_engine


def build_xlsx(run: Dict[str, Any], attributes: List[Dict[str, Any]], config: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    label = {p["slug"]: p["label"] for p in config["products"]}
    cname = {c["code"]: c["label"] for c in config["countries"]}
    akey = lambda a: a["code"] + ("" if not a.get("sub") else f"#{a['no']}")
    hfont, hfill = Font(bold=True, color="FFFFFF", name="Arial"), PatternFill("solid", fgColor="8A5A00")

    wb = Workbook()
    ws = wb.active; ws.title = "Summary"
    ws.append(["[honeyComb · Google Shopping Monitoring]", run["week"], run["at"], f"source={run.get('source')}"])
    ws.append([f"*판정: position 1 = 1순위 · position ≤ {config['top_n']} = 상단 노출(둘째 줄) · gl/hl 지정, 비로그인 데스크톱 기준"])
    ws.append([])
    ws.append(["Country", "Product", "Keyword", "Keyword Type", "Status", "Position", "S.com (O/X)", "1위 판매처",
               "Entered Attribute", "Missed Attribute"])
    for c in ws[4]:
        c.font, c.fill = hfont, hfill
    for cell in sorted(run["cells"], key=lambda x: (x["product"], x["country"], x["keyword_type"])):
        a = cell.get("attrs") or {}
        ent = [x["name"] for x in attributes if a.get(akey(x)) == "entered"]
        mis = [x["name"] for x in attributes if a.get(akey(x)) == "missed"]
        ws.append([cname.get(cell["country"], cell["country"]), label.get(cell["product"], cell["product"]), cell["keyword"],
                   f"{cell.get('keyword_subtype') or ''} {'제품명' if cell.get('keyword_type') == 'brand' else '자연어'}".strip(),
                   cell["status"], cell.get("position"), cell.get("scom_exposed") or "", cell.get("first_store") or "",
                   f"Entered ({len(ent)}): " + ", ".join(ent) if ent else "", f"Missed ({len(mis)}): " + ", ".join(mis) if mis else ""])
    for col, w in zip("ABCDEFGHIJ", (14, 22, 34, 10, 9, 10, 22, 60, 60, 24)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Shopping Check")
    ws2.append(["Country", "Product", "Keyword", "Attribute Name", "Observe", "Entered (O/X)", "Missed (O/X)"])
    for c in ws2[1]:
        c.font, c.fill = hfont, hfill
    for cell in run["cells"]:
        a = cell.get("attrs") or {}
        if not a:
            continue
        for x in attributes:
            k = akey(x); v = a.get(k, "na")
            ws2.append([cname.get(cell["country"], cell["country"]), label.get(cell["product"], cell["product"]), cell["keyword"],
                        x["name"], x["observe"], "O" if v == "entered" else "X" if v == "missed" else "-",
                        "O" if v == "missed" else "X" if v == "entered" else "-"])
    for col, w in zip("ABCDEFG", (14, 22, 34, 30, 8, 12, 12)):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("Appendix")
    ws3.append(["No.", "Attribute Category", "Attribute Name", "Attribute Code", "Feeding", "Observe(card/detail/feed)", "Conversational"])
    for c in ws3[1]:
        c.font, c.fill = hfill and hfont, hfill
    for x in attributes:
        ws3.append([x["no"], x["category"], x["name"], x["code"], x.get("feeding", ""), x["observe"], "Y" if x.get("conversational") else ""])
    for col, w in zip("ABCDEFG", (6, 32, 32, 30, 18, 22, 14)):
        ws3.column_dimensions[col].width = w
    wb.create_sheet("Screenshot").append(["(실수집 provider 연결 후 증빙 이미지/HTML 참조가 여기에 들어갑니다)"])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
